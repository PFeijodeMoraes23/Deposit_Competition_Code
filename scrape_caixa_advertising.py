"""
scrape_caixa_advertising.py
===========================
CAIXA ECONOMICA FEDERAL's own published advertising outlays, monthly from January 2013, and its
published sponsorship contracts (a separate robustness series), 2019 onward.

Caixa files with no securities regulator, so its transparency page is the only pre-2025
advertising series for it. The documents are titled "Custos CAIXA <mes>" or "Despesas com
Publicidade CAIXA - <mes> (R$)". No footnote says whether the amounts are billed, committed or
paid, and no agency-commission line is shown; April 2019 prints "Nao existem pagamentos
registrados no periodo". Every amount is therefore tagged basis = "custos_undocumented".

Sources (www.caixa.gov.br, SharePoint list API; plain clients loop on a 302 without a cookie jar)
------------------------------------------------------------------------------------------------
  _api/web/lists/GetByTitle('despesas-publicidade')/Items   one item per year, links in ColExtra2
      "<Mes>"             the PDF (2019: a ZIP holding the summary PDF and a supplier PDF)
      "<Mes> - Editavel"  from April 2020, a ZIP holding an XLSX (some months .xls, from November
                          2025 .xltx templates, May 2025 a PDF, January 2022 a PDF and an XLSX)
      "Todos os meses"    2010-2016 year ZIPs; every member is byte-identical to the monthly PDF
                          (checked 2026-09-14), so they are not downloaded
  _api/web/lists/GetByTitle('patrocinio')/Items             one item per year 2019-2026, a ZIP with
                          "Contratados" (general programme) and, from 2023, "TransparenciaPOCC"
                          (Programa de Ocupacao dos espacos da CAIXA Cultural; the 2023 one is an
                          HTML table saved as .xls)
Downloads are cached under paths.AWARENESS_RAW / "caixa"; --refresh downloads again.

Document layouts
----------------
  2013-01 .. 2014-06, 2016-04, 2016-05  summary page is an image scan (no text layer; only the
                          supplier list pages carry text). Status unreadable_scan here, never
                          estimated; scrape_caixa_scan_advertising.py reads them by OCR.
  2014-07 .. 2015-04, 2015-07 .. 2018-12  unruled layout: category labels on the left, one column
                          per agency. Parsed from word positions: each amount is attached to the
                          nearest label line (flag ambiguous_rows when that line is more than
                          4.5pt away), labels wrapped onto a value-less line are joined. In the
                          amount columns a lone thousands-grouped integer ("14.250") is a
                          whole-number cell (rule below), a US-format amount ("4,641,371.16",
                          April 2015) is read on a TOTAL row only, and any other word with digits
                          there is noted as not read, never dropped silently.
  2015-05, 2015-06, 2019-01 .. 2026-06  ruled table. Parsed from the table cells; merged cells
                          carry their label down; vertical (rotated) labels are read bottom-to-top
                          from the characters; a table split across pages continues.
  Spreadsheets            the first sheet with an agency header row; merged ranges carry labels.
Only the summary table is read (section > category > subcategory x agency). The supplier lists
(names and CNPJs, no amounts) are not.

Rules
-----
  - Amounts are reais as published. Blank, "-" and 0,00 cells are not lines (the 2014-15 layout
    prints 0,00 in every empty cell, later layouts leave them blank).
  - A whole-number cell (a thousands-grouped integer with no decimals, "14.250") is read as money
    only when Caixa's printed totals include it: its agency's printed TOTAL or the TOTAL GERAL
    closes with the cell and does not close without it. Otherwise it is not read, and the month's
    notes give its page, row and column. May 2017's NOVA/SB JORNAL "14.250" is read (both totals
    close only with it); October 2017's NOVA/SB BANNERS DIGITAIS "138.500", printed beside
    HEADS's 138.500,00 on the same row, is not (both totals close without it).
  - A US-format amount ("4,641,371.16") is read on a total row only, and noted; April 2015
    prints one agency's TOTAL so. Anywhere else it is noted as not read.
  - Labels are kept as published, with whitespace collapsed. Nothing is reclassified: events,
    fees (cachês), research, copyright and similar rows stay in, each under its own label.
  - One source per month (parse_method): the spreadsheet where one parses (xlsx, or template for
    .xltx), else the PDF found in the spreadsheet slot (pdf_in_xlsx_slot), else the PDF
    (text_pdf). Where both a spreadsheet and a PDF parse, they are compared.
  - stated_total_brl is the grand total ("TOTAL GERAL") printed in the chosen document; when that
    document has none (a formula without a stored value, or no row), the other document's is used
    and stated_total_source says which. Per-agency "TOTAL" rows are checked as a flag. Two cases
    validate against a total that is not a printed TOTAL GERAL: April 2019 prints "Nao existem
    pagamentos registrados no periodo" and no table, so its total is a synthetic 0.0 taken from
    that sentence; in July to November 2014 the TOTAL GERAL row carries no amount in the text
    layer, so their total is the sum of the per-agency TOTAL rows (stated_total_source
    "text_pdf:agency_total_row"), and there the total check and the agency checks test the same
    printed figures.

Outputs (paths.AWARENESS_PROC)
------------------------------
  caixa_advertising_lines.{parquet,csv}     one row per published nonzero cell
  caixa_advertising_monthly.{parquet,csv}   one row per month 2013-01 .. latest published month
  caixa_sponsorship_contracts.{parquet,csv} one row per contract
All carry bank_key = "caixa", cnpj8 = "00360305" and panel_key, the prudential conglomerate code
established from the IF.data lists (CnpjInstituicaoLider 00360305) and confirmed against the
conglomerate rows of cosif_advertising_monthly.parquet.

Monthly flags
-------------
Every check below compares integer cents: the gap is the components' sum minus the printed
total (spreadsheet minus PDF for pdf_xlsx_agree). A gap of 0 is exact, a gap of at most
MAX_GAP_CENTS (5 centavos) either way is accepted and flagged, anything larger fails.
  validation_status        validated_exact (every check 0) | validated_within_5_centavos (every
                           check within 5 centavos, one or more not 0) | not_validated_document
                           (some check above 5 centavos, the month is on the reviewed list with
                           exactly these gaps, and its image evidence is complete; below) |
                           not_validated_parser (some check above 5 centavos, and anything else:
                           the parser missed an amount, the evidence is incomplete or refused,
                           the image could not be read, or the month looks like the document's
                           but no person has reviewed it with these gaps) | no_stated_total (no
                           printed total was read, which includes the months with no parsed
                           document; see status)
  validated                True for validated_exact and validated_within_5_centavos only; a
                           document-inconsistent month is never validated
  validation_reason        for a month not validated, why it is the document's or the parser's;
                           a month the evidence puts on the document's side but that is not on
                           the reviewed list with these gaps starts "unreviewed document
                           candidate"
  total_gap_cents          lines minus the stated total, in cents (NA without a stated total)
  max_gap_cents            largest absolute gap over the month's checks (NA when none ran)
  tolerance_flagged        validation_status is validated_within_5_centavos, as in
                           scrape_caixa_scan_advertising.py
  tolerance_detail         every check that fell within the allowance and its gap, e.g.
                           "total -3", also in a month that fails another check
  gap_detail               every check with a nonzero gap and that gap in cents; where a
                           spreadsheet-PDF disagreement is explained by the PDF's cells not adding
                           up to its own agency totals, also those gaps ("pdf_cells_vs_own_total
                           <agency>"), which do not enter the status
  image_*                  the image check of a month that fails (below): image_check (ran, or
                           why not), image_pages, image_amounts (nonzero amounts OCR read),
                           image_unparsed and image_unparsed_brl (amounts the image shows that
                           the parser did not use, with page and nearest agency header),
                           image_missing and image_missing_brl (parser amounts OCR did not find),
                           image_totals_missing (printed totals OCR did not find),
                           image_unreadable (OCR text that could not be read as amounts: a run
                           with a decimal comma, or of 4 or more digits with a dot or comma),
                           image_outside_amounts (Brazilian-format amounts where the image is not
                           read: the text layer of the pages not imaged, and the text layer and
                           OCR above the header row), image_lines_on_total_rows (parser lines
                           whose amount the image shows only on a row labelled TOTAL)
  printed_totals_conflict  for a month that fails: the printed agency totals do not add up to the
                           printed TOTAL GERAL, in a way that accounts for every failing check
  total_matches            the total check passed (gap within 5 centavos)
  agency_totals_match      lines per agency sum to the printed per-agency TOTAL row
  pdf_xlsx_agree           spreadsheet and PDF agree on every agency sum (NaN if not both parsed)
  pdf_xlsx_lines_match     share of spreadsheet (agency, amount) lines also found in the PDF
  repeated_agency_block    agencies whose whole block (labels and amounts) equals the previous month
  title_month_mismatch     the document title names another month or year
  ambiguous_rows           unruled layout: amounts attached to a label line more than 4.5pt away
  no_payments_stated       the document says no payments were registered in the period
  pdf_slot_broken          the PDF link returns a web page instead of a PDF

Whose gap is it (months that fail only; explain_failure)
--------------------------------------------------------
The pages of the chosen document that hold the summary table (those the parser read it from, and
any page without a text layer) are rendered with PyMuPDF at 200 dpi and read with RapidOCR; for a
spreadsheet month, the pages of the PDF published beside it. Every nonzero amount OCR finds is set
against the amounts the parser used (as a multiset of amounts) and the printed totals. OCR's known
misreads of an amount are read as the amount: the decimal comma read as a dot ("26.453.70"), a
space before the cents ("26.453, 70"), US format ("26,453.70"). The steps, in order:
  1. An amount the image shows that the parser did not use: parser failure.
  2. Incomplete evidence: parser failure. The evidence is complete only if OCR read at least one
     amount, found every parser amount and every printed total, left no run of digits unreadable,
     and no Brazilian-format amount lies where the image is not read (the text layer of the pages
     not imaged, i.e. supplier lists and any page after the TOTAL GERAL, which the parser never
     opens; and the text layer and OCR above the header row). Printed totals that disagree with
     each other do not stand in for missing evidence: the parser read those totals itself, and a
     misread total produces exactly such a disagreement.
  3. Refusals, a parser failure whatever else holds: the total gap equals, within 5 centavos, a
     sum of amounts printed on total rows (a total row read as data doubles those amounts); or a
     parser line's amount is one the image shows only on a row labelled TOTAL.
  4. With complete evidence the lines add up to what the document prints. A TOTAL GERAL that does
     not close, or (when only per-agency totals fail, which a comparison of amounts cannot settle
     because it does not see columns) printed totals that disagree with each other in a way that
     accounts for every failing check (April 2017), make the month a document candidate. An
     agency-only failure without such a disagreement stays a parser failure.
  5. A document candidate is not_validated_document only if REVIEWED_DOCUMENT_MONTHS lists it with
     exactly its failing checks and gaps in cents. Otherwise it is not_validated_parser, reason
     "unreviewed document candidate", and counts against the gate until a person reviews it. The
     list is what keeps a parser that counts amounts it should not (a total row read as data, a
     block read twice) from passing: its months also show every image amount used.

Reviewing a month and adding it to REVIEWED_DOCUMENT_MONTHS
-----------------------------------------------------------
  1. Take the month's gap_detail and validation_reason from the log ("unreviewed document
     candidate") and the pages in image_pages.
  2. Render those pages (200 dpi) and check by eye that every amount printed in the table is a
     line or a printed total, and that the failing agency's cells, summed by hand, differ from its
     printed total by the logged gap.
  3. Search every page of the document, supplier lists included, for the gap amount (OCR where
     there is no text layer); it must not be printed anywhere the parser does not read.
  4. Add period: dict(gaps={check: gap in cents, for every check above 5 centavos, exactly as
     gap_detail prints them}, reason=<one line>, reviewed=<date>) to the list and a line below.
A parse change that moves any gap by a centavo, or adds or removes a failing check, sends the
month back to not_validated_parser until it is reviewed again.

Months whose printed tables do not add up (REVIEWED_DOCUMENT_MONTHS, as of the 2026-09 files)
---------------------------------------------------------------------------------------------
The cells are kept as printed; each gap is lines minus the printed figure. Each month was read
by eye and checked by an independent OCR of every page.
  2015-07  TOTAL GERAL +295,004.21: ARTPLAN +121,801.26, exactly its BANNERS/GRANDES FORMATOS
           cell (printed again, and counted, in August); HEADS -138,512.00; NOVA/SB +311,714.98.
  2015-08  +225,798.10: NOVA/SB +227,048.10, exactly one of its two cells printed 227.048,10;
           ARTPLAN -1,250.00.
  2016-03  -30,179.88, all HEADS.
  2016-06  -26,453.70, all PROPEG; the agency totals add up to the TOTAL GERAL.
  2016-07  -43,200.00, all ARTPLAN. Separately, the FOTO cell 14.398,31 is printed under
           NOVA/SB, whose total is exactly that much below its cells while PROPEG's is that much
           above: the totals seem to count it under PROPEG (an inference).
  2017-04  lines match the TOTAL GERAL; NOVA/SB's printed total is 10.00 below its cells, so the
           agency totals sum to 10.00 less than the TOTAL GERAL.
  2023-07  -40,392.00, all BINDER (spreadsheet and PDF alike; the totals are typed values).
  2024-05  the TOTAL row adds up to the cells and the TOTAL GERAL does not: -264,748.96 in May,
  2024-07  -175,357.83 in July (typed values in the spreadsheet, not formulas).
Neither 26.453,70 (June 2016) nor 43.200,00 (July 2016) appears anywhere in the OCR of every page
of those documents, supplier lists included.

Validation (a hard failure aborts before writing)
-------------------------------------------------
  1. Every month from 2013-01 to the latest published month has a status.
  2. Parser health. Of the months with a stated total, validated months plus
     not_validated_document months must be at least 95%, i.e. parser failures at most 5%. The
     gate exists to catch a broken parser; a reviewed document inconsistency is logged with its
     evidence and never validated, but does not count against the parser, while an unreviewed
     document candidate does. The count of months per validation_status is logged, and so is
     every month not validated.
  3. Spreadsheet and PDF agree on every agency sum within 5 centavos wherever both parse.
  4. No negative amount.
  5. April, May and June 2020 reproduce R$ 2.4m, 17.2m and 35.6m.
  6. The panel key is unique and confirmed by COSIF.
  7. Sponsorship: every contract has a positive amount and a date.
Printed, not gated: 2025 ratios of Caixa's quarterly, half-year and annual sums to the COSIF
conglomerate flows adv and all3 (cosif_advertising_quarterly.parquet, entity C0080738).

Usage
-----
  python scrape_caixa_advertising.py
  python scrape_caixa_advertising.py --refresh           # download the lists and files again
  python scrape_caixa_advertising.py --out-dir <dir>     # write somewhere else (smoke runs)
"""

from __future__ import annotations

import argparse
import datetime as dt
import html
import io
import json
import logging
import re
import time
import unicodedata
import zipfile
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter
import pdfplumber
import requests
import xlrd
from bs4 import BeautifulSoup

from utils import paths
from utils.disclosure_common import http_get

logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(levelname)-7s  %(message)s",
                    datefmt="%H:%M:%S")
logging.getLogger("pdfminer").setLevel(logging.ERROR)
log = logging.getLogger(__name__)

RAW_DIR = paths.AWARENESS_RAW / "caixa"
OUT_DIR = paths.AWARENESS_PROC
BASE = "https://www.caixa.gov.br"
LIST_API = BASE + "/_api/web/lists/GetByTitle('{title}')/Items/?$top=5000"
AD_LIST, SPONSOR_LIST = "despesas-publicidade", "patrocinio"

BANK_KEY, CNPJ8 = "caixa", "00360305"
BASIS_AD, BASIS_SPONSOR = "custos_undocumented", "contracted"
FIRST_MONTH = "201301"

AGENCIES = {"ARTPLAN", "BORGHI", "HEADS", "NOVA/SB", "NOVASB", "PROPEG", "PROPER", "BINDER",
            "CALIA"}
AGENCY_CANON = {"NOVASB": "NOVA/SB", "PROPER": "PROPEG"}   # October 2020 header misprint
MONTHS_PT = {"janeiro": 1, "fevereiro": 2, "marco": 3, "abril": 4, "maio": 5, "junho": 6,
             "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11,
             "dezembro": 12}
MONEY_RE = re.compile(r"^-?\d{1,3}(?:\.\d{3})*,\d{1,4}$|^-?\d+,\d{1,4}$")
# A whole-number cell: thousands dots and no decimal comma. The corpus has two. May 2017 prints
# NOVA/SB JORNAL as "14.250", and its agency total and TOTAL GERAL close only with it as
# 14,250.00. October 2017 prints NOVA/SB BANNERS DIGITAIS as "138.500" beside HEADS's
# 138.500,00 on the same row, and both totals close without it. So such a cell is read only when
# the printed totals include it (resolve_whole_numbers), which also keeps a stray CNPJ root
# ("11.222.333") or a percentage whose '%' sits apart from being read as money. No leading zero,
# so a CNPJ root such as "00.360.305" is never even a candidate.
INT_AMOUNT_RE = re.compile(r"^[1-9]\d{0,2}(?:\.\d{3})+$")
# US format, read on a total row only: April 2015 prints one agency's TOTAL as "4,641,371.16".
US_AMOUNT_RE = re.compile(r"^\d{1,3}(?:,\d{3})+\.\d{2}$")
# Recognising a total row printed without a label (parse_pdf_ruled) only; see within_tol.
TOL_ABS, TOL_REL = 1.0, 0.001
# The documents print cents, so components and the total printed for them agree to the cent or
# disagree, whether the fault is in the document's arithmetic or in the parse. A disagreement of
# up to this many centavos is accepted, never silently: the month records it (tolerance_flagged).
# A relative tolerance would accept tens of thousands of reais on a R$40m month.
MAX_GAP_CENTS = 5
MIN_MATCH_SHARE = 0.95
# Months a person has reviewed and found to be Caixa's own inconsistency (module docstring:
# "Reviewing a month"). period -> every check above MAX_GAP_CENTS with its exact gap in cents
# (lines minus the printed figure), a one-line reason, and the date of the review. The evidence
# alone cannot tell a Caixa typo from a parser that counts amounts it should not (a total row read
# as data leaves every image amount used and the TOTAL GERAL unclosed), so no month is the
# document's until someone has looked at it with exactly these gaps.
REVIEWED_DOCUMENT_MONTHS: dict[str, dict] = {
    "201507": dict(gaps={"total": 29500421, "agency_total ARTPLAN": 12180126,
                         "agency_total HEADS": -13851200, "agency_total NOVA/SB": 31171498},
                   reason="TOTAL GERAL does not close; ARTPLAN's gap is exactly its BANNERS/"
                          "GRANDES FORMATOS cell", reviewed="2026-09-28"),
    "201508": dict(gaps={"total": 22579810, "agency_total ARTPLAN": -125000,
                         "agency_total NOVA/SB": 22704810},
                   reason="TOTAL GERAL does not close; NOVA/SB's gap is exactly one of its two "
                          "cells printed 227.048,10", reviewed="2026-09-28"),
    "201603": dict(gaps={"total": -3017988, "agency_total HEADS": -3017988},
                   reason="HEADS's printed total exceeds its cells; the gap is printed nowhere",
                   reviewed="2026-09-28"),
    "201606": dict(gaps={"total": -2645370, "agency_total PROPEG": -2645370},
                   reason="PROPEG's printed total exceeds its cells; 26.453,70 is printed nowhere "
                          "in the document", reviewed="2026-09-28"),
    "201607": dict(gaps={"total": -4320000, "agency_total ARTPLAN": -4320000,
                         "agency_total NOVA/SB": 1439831, "agency_total PROPEG": -1439831},
                   reason="ARTPLAN's printed total exceeds its cells (43.200,00 printed nowhere); "
                          "the FOTO cell 14.398,31 printed under NOVA/SB is totalled under PROPEG",
                   reviewed="2026-09-28"),
    "201704": dict(gaps={"agency_total NOVA/SB": 1000},
                   reason="lines match the TOTAL GERAL; NOVA/SB's printed total is 10.00 below its "
                          "cells and the agency totals sum to 10.00 less than the TOTAL GERAL",
                   reviewed="2026-09-28"),
    "202307": dict(gaps={"total": -4039200, "agency_total BINDER": -4039200},
                   reason="BINDER's typed total exceeds its cells, spreadsheet and PDF alike",
                   reviewed="2026-09-28"),
    "202405": dict(gaps={"total": -26474896},
                   reason="the TOTAL row equals the cells; the typed TOTAL GERAL is higher",
                   reviewed="2026-09-28"),
    "202407": dict(gaps={"total": -17535783},
                   reason="the TOTAL row equals the cells; the typed TOTAL GERAL is higher",
                   reviewed="2026-09-28"),
}
AMBIGUOUS_PT = 4.5
KNOWN_TOTALS_M = {"202004": 2.4, "202005": 17.2, "202006": 35.6}


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------
def fold(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s))
    return "".join(c for c in s if not unicodedata.combining(c)).lower()


def clean(s) -> str:
    if s is None:
        return ""
    return re.sub(r"[\s​\xa0]+", " ", str(s)).strip()


def agency_name(s) -> str | None:
    t = re.sub(r"\s+", "", clean(s)).upper()
    return AGENCY_CANON.get(t, t) if t in AGENCIES else None


def parse_brl(s) -> float | None:
    """Brazilian-format amount -> float; None for a blank or dash cell; ValueError otherwise."""
    if s is None:
        return None
    if isinstance(s, (int, float)) and not isinstance(s, bool):
        return None if pd.isna(s) else float(s)
    t = re.sub(r"[\s​\xa0]+", "", str(s)).replace("R$", "")
    if t in ("", "-", "–", "—"):
        return None
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()")
    if re.fullmatch(r"-?[\d.]*,\d{1,4}", t) or re.fullmatch(r"-?\d{1,3}(\.\d{3})+", t):
        v = float(t.replace(".", "").replace(",", "."))
    elif re.fullmatch(r"-?\d+(\.\d+)?", t):
        v = float(t)
    else:
        raise ValueError(f"not an amount: {s!r}")
    return -v if neg else v


def title_period(text: str) -> str | None:
    """YYYYMM named in a document title such as 'Custos CAIXA MAIO/2015'."""
    f = fold(text)
    m = re.search(r"(janeiro|fevereiro|marco|abril|maio|junho|julho|agosto|setembro|outubro|"
                  r"novembro|novemro|dezembro)\s*(?:de\s*)?[/-]?\s*(\d{4})", f)
    if not m:
        return None
    month = MONTHS_PT.get(m.group(1), 11 if m.group(1) == "novemro" else None)
    return f"{m.group(2)}{month:02d}"


# ---------------------------------------------------------------------------
# Downloading
# ---------------------------------------------------------------------------
class Fetcher:
    """utils.disclosure_common.http_get first. www.caixa.gov.br answers a client without cookies
    with a 302 loop (and file URLs with 429), so on failure a requests session takes over; it is
    primed on the list API, which sets the cookie."""

    def __init__(self) -> None:
        self.session: requests.Session | None = None

    def _session(self) -> requests.Session:
        if self.session is None:
            self.session = requests.Session()
            self.session.headers["User-Agent"] = "Mozilla/5.0 (research; deposit-competition)"
            self.session.get(LIST_API.format(title=AD_LIST), timeout=90,
                             headers={"Accept": "application/json;odata=verbose"})
        return self.session

    def get(self, url: str, accept: str = "*/*") -> bytes:
        if self.session is None:
            status, body = http_get(url, binary=True, retries=1, extra_headers={"Accept": accept})
            if status == 200 and body is not None:
                return body
            log.info("http_get gave status %s for %s; using a cookie session", status, url)
        s = self._session()
        for attempt in range(1, 6):
            try:
                r = s.get(url, headers={"Accept": accept}, timeout=120)
                if r.status_code == 200:
                    time.sleep(0.3)
                    return r.content
                if r.status_code in (403, 404, 410):
                    raise FileNotFoundError(f"HTTP {r.status_code} for {url}")
                wait = float(r.headers.get("Retry-After") or 5 * attempt)
                log.warning("HTTP %s on %s (attempt %d/5), waiting %.0fs", r.status_code, url,
                            attempt, wait)
            except requests.RequestException as e:
                wait = 3 * attempt
                log.warning("GET failed %s on %s (attempt %d/5)", type(e).__name__, url, attempt)
            time.sleep(wait)
        raise RuntimeError(f"could not download {url}")


def list_items(fetcher: Fetcher, title: str, refresh: bool) -> list[dict]:
    cache = RAW_DIR / f"list_{title}.json"
    if refresh or not cache.exists():
        body = fetcher.get(LIST_API.format(title=title), accept="application/json;odata=verbose")
        json.loads(body.decode("utf-8"))
        cache.parent.mkdir(parents=True, exist_ok=True)
        cache.write_bytes(body)
    return json.loads(cache.read_text(encoding="utf-8"))["d"]["results"]


def link_url(href: str) -> str:
    href = html.unescape(href).strip()
    return BASE + href if href.startswith("/") else re.sub(r"^http://", "https://", href)


def item_links(item: dict) -> list[tuple[str, str]]:
    soup = BeautifulSoup(item.get("ColExtra2") or "", "lxml")
    return [(clean(a.get_text(" ")).strip("• "), link_url(a["href"]))
            for a in soup.find_all("a", href=True)]


def advertising_documents(items: list[dict]) -> list[dict]:
    docs = []
    for it in items:
        m = re.search(r"(\d{4})", it.get("Title") or "")
        if not m:
            raise ValueError(f"list item without a year: {it.get('Title')!r}")
        year = int(m.group(1))
        for label, url in item_links(it):
            f = fold(label)
            if "todos os meses" in f:
                continue
            month = next((n for k, n in MONTHS_PT.items() if f.startswith(k)), None)
            if month is None:
                raise ValueError(f"unrecognised link label {label!r} in {it['Title']}")
            period = f"{year}{month:02d}"
            if period < FIRST_MONTH:
                continue
            docs.append(dict(period=period, slot="editable" if "editavel" in f else "pdf",
                             label=label, url=url))
    dup = pd.DataFrame(docs).duplicated(["period", "slot"])
    if dup.any():
        raise ValueError(f"two links for the same month and slot: "
                         f"{pd.DataFrame(docs)[dup].to_dict('records')}")
    return docs


def sponsorship_documents(items: list[dict]) -> list[dict]:
    docs = []
    for it in items:
        m = re.search(r"(\d{4})", it.get("Title") or "")
        if not m or "patroc" not in fold(it.get("Title") or ""):
            raise ValueError(f"unexpected sponsorship list item {it.get('Title')!r}")
        for _, url in item_links(it):
            docs.append(dict(year=int(m.group(1)), url=url))
    return docs


def download(fetcher: Fetcher, url: str, dest: Path, refresh: bool) -> Path:
    if dest.exists() and dest.stat().st_size > 0 and not refresh:
        return dest
    dest.parent.mkdir(parents=True, exist_ok=True)
    body = fetcher.get(url)
    tmp = dest.with_name(dest.name + ".part")
    tmp.write_bytes(body)
    tmp.replace(dest)
    log.info("  downloaded %s (%d bytes)", dest.name, len(body))
    return dest


def ad_path(doc: dict) -> Path:
    name = Path(requests.utils.urlparse(doc["url"]).path).name
    return RAW_DIR / "despesas" / doc["period"][:4] / f"{doc['period']}_{doc['slot']}__{name}"


def sponsor_path(doc: dict) -> Path:
    name = Path(requests.utils.urlparse(doc["url"]).path).name
    return RAW_DIR / "patrocinio" / f"{doc['year']}__{name}"


# ---------------------------------------------------------------------------
# Document containers
# ---------------------------------------------------------------------------
@dataclass
class Parsed:
    method: str                       # text_pdf | xlsx | template | pdf_in_xlsx_slot
    member: str                       # file name inside the downloaded file
    lines: list[dict] = field(default_factory=list)
    stated_total: float | None = None
    agency_totals: dict[str, float] = field(default_factory=dict)
    # Every amount printed on a total row, as printed (a grand total printed per agency is kept
    # cell by cell here, while stated_total is its sum). The image check needs them to tell a
    # printed total from a cell the parser missed.
    total_row_amounts: list[float] = field(default_factory=list)
    table_pages: list[int] = field(default_factory=list)   # 1-based pages the summary table read
    agencies: list[str] = field(default_factory=list)
    status: str = "parsed"            # parsed | unreadable_scan
    title_period: str | None = None
    ambiguous_rows: int = 0
    no_payments_stated: bool = False
    notes: list[str] = field(default_factory=list)

    @property
    def total(self) -> float:
        return float(sum(x["amount_brl"] for x in self.lines))

    def agency_sums(self) -> dict[str, float]:
        out = {a: 0.0 for a in self.agencies}
        for x in self.lines:
            out[x["agency"]] = out.get(x["agency"], 0.0) + x["amount_brl"]
        return out

    def total_cents(self) -> int:
        return cents(x["amount_brl"] for x in self.lines)

    def agency_cents(self) -> dict[str, int]:
        names = dict.fromkeys(self.agencies) | dict.fromkeys(x["agency"] for x in self.lines)
        return {a: cents(x["amount_brl"] for x in self.lines if x["agency"] == a) for a in names}


def cents(amounts) -> int:
    """Sum of amounts in reais, as integer cents. Each amount enters as the decimal it was read
    as and the sum is exact; only the sum is rounded, half away from zero. A spreadsheet cell can
    hold more than two decimals and a PDF cell up to four (MONEY_RE), while a total is computed
    from the unrounded cells: rounding each cell first could move the sum by up to half a cent
    per cell and invent a gap the document does not have."""
    s = sum((Decimal(repr(float(v))) for v in amounts), Decimal(0))
    return int((s * 100).to_integral_value(rounding=ROUND_HALF_UP))


def unpack(path: Path) -> list[tuple[str, bytes]]:
    """(member name, bytes) for a downloaded file: ZIP members, or the file itself."""
    b = path.read_bytes()
    if b[:2] == b"PK" and path.suffix.lower() == ".zip":
        with zipfile.ZipFile(io.BytesIO(b)) as z:
            return [(m, z.read(m)) for m in z.namelist() if not m.endswith("/")]
    return [(path.name, b)]


def kind_of(name: str, b: bytes) -> str:
    low = name.lower()
    if b[:4] == b"%PDF":
        return "pdf"
    if b[:2] == b"PK" and low.endswith((".xlsx", ".xltx", ".xlsm")):
        return "template" if low.endswith(".xltx") else "xlsx"
    if b[:8] == bytes.fromhex("d0cf11e0a1b11ae1"):
        return "xls"
    if b.lstrip()[:1] == b"<":
        return "html"
    return "unknown"


# ---------------------------------------------------------------------------
# Resolving label hierarchies (tables and spreadsheets)
# ---------------------------------------------------------------------------
TOTAL_GERAL_RE = re.compile(r"^total\s*geral", re.I)


def row_kind(labels: list[str]) -> str:
    texts = [fold(t) for t in labels if t]
    if any(TOTAL_GERAL_RE.match(t) for t in texts):
        return "grand_total"
    if any(t.strip(" :") == "total" or t.startswith("total por") for t in texts):
        return "agency_total"
    return "data"


def resolve_sections(rows: list[dict], n_section_cols: int) -> None:
    """Fill row['section'] from the section columns. A column whose vertical label was split
    into several stacked unmerged cells (e.g. January 2019) is joined bottom-to-top and each row
    takes the nearest such label; otherwise merged cells carry down."""
    data_idx = [i for i, r in enumerate(rows) if r["kind"] == "data"]
    parts = []
    for c in range(n_section_cols):
        cells = [rows[i]["sec_raw"][c] for i in data_idx]
        texts = [x["text"] if x else None for x in cells]
        stacked = any(texts[k] and texts[k + 1] for k in range(len(texts) - 1))
        values = [""] * len(data_idx)
        if stacked:
            runs, k = [], 0
            while k < len(texts):
                if texts[k]:
                    j = k
                    while j + 1 < len(texts) and texts[j + 1]:
                        j += 1
                    pieces = [cells[t] for t in range(k, j + 1)]
                    vertical = any(p["rotated"] for p in pieces) or all(
                        len(p["text"]) <= 3 for p in pieces)
                    word = ("".join(p["text"].replace(" ", "") for p in reversed(pieces))
                            if vertical else " ".join(p["text"] for p in pieces))
                    runs.append(((k + j) / 2, word))
                    k = j + 1
                else:
                    k += 1
            for pos in range(len(data_idx)):
                if runs:
                    values[pos] = min(runs, key=lambda r: (abs(r[0] - pos), r[0]))[1]
        else:
            current = ""
            for pos, cell in enumerate(cells):
                if cell is None:
                    values[pos] = current
                else:
                    current = cell["text"]
                    values[pos] = current
        parts.append(values)
    for pos, i in enumerate(data_idx):
        rows[i]["section"] = " > ".join(p[pos] for p in parts if p[pos])


def cell_position(r: dict, agency: str) -> str:
    """Where a cell sits, for a note: page or sheet, row label, agency column, and the cell's
    coordinates or address where the parser kept them."""
    label = (" / ".join(t for t in (r.get("group"), r.get("item")) if t)
             or " ".join(t for t in r.get("label_texts", []) if t) or "(no label)")
    at = r.get("pos", {}).get(agency)
    return f"{r['where']}, row {label!r}, column {agency}" + (f" ({at})" if at else "")


def total_cell_amount(raw, r: dict, agency: str, parsed: Parsed) -> float | None:
    """An amount printed on a total row: Brazilian format, or US format, which is noted."""
    if isinstance(raw, str) and US_AMOUNT_RE.fullmatch(clean(raw)):
        t = clean(raw)
        parsed.notes.append(f"US-format amount on a total row read as {t.replace(',', '')}: "
                            f"{t!r} at {cell_position(r, agency)}")
        return float(t.replace(",", ""))
    return parse_brl(raw)


def resolve_whole_numbers(whole: list[tuple[dict, str, str, int]], parsed: Parsed) -> None:
    """Read each whole-number cell as money only when Caixa's printed totals include it: its
    agency's printed TOTAL or the TOTAL GERAL closes (within MAX_GAP_CENTS) with the cell and does
    not close without it. The sums without it leave out every whole-number cell of the document,
    so several such cells that the totals need only together are none of them read; the month
    then fails, and the image check finds them unparsed. A cell read is inserted where its row
    put it among the lines (the index kept with it)."""
    base_all = parsed.total_cents()
    base_agency = parsed.agency_cents()
    stated = None if parsed.stated_total is None else cents([parsed.stated_total])

    def closes(sum_cents: int, printed: int) -> bool:
        return abs(sum_cents - printed) <= MAX_GAP_CENTS

    accepted = []
    for r, agency, text, at in whole:
        c = cents([parse_brl(text)])
        why = []
        if agency in parsed.agency_totals:
            own, base = cents([parsed.agency_totals[agency]]), base_agency.get(agency, 0)
            if closes(base + c, own) and not closes(base, own):
                why.append(f"{agency}'s printed TOTAL")
        if stated is not None and closes(base_all + c, stated) and not closes(base_all, stated):
            why.append("the TOTAL GERAL")
        where = cell_position(r, agency)
        if why:
            accepted.append((at, dict(agency=agency, section=r["section"], category=r["group"],
                                      subcategory=r["item"], amount_brl=c / 100,
                                      page_or_sheet=r["where"])))
            verb = "close" if len(why) > 1 else "closes"
            parsed.notes.append(f"whole-number cell {text!r} at {where} read as {brl_text(c)}: "
                                f"{' and '.join(why)} {verb} with it and not without it")
        else:
            parsed.notes.append(f"whole-number cell {text!r} at {where} not read: Caixa's printed "
                                f"totals do not include it")
    for _, (at, line) in sorted(enumerate(accepted), key=lambda x: (-x[1][0], -x[0])):
        parsed.lines.insert(at, line)


def lines_from_rows(rows: list[dict], agencies: list[str], parsed: Parsed) -> None:
    """rows: dicts with kind, section, group, item, values {agency: raw}, where, and optionally
    pos {agency: coordinates or cell address} for notes."""
    whole: list[tuple[dict, str, str, int]] = []
    for r in rows:
        if r["kind"] == "grand_total":
            vals = [total_cell_amount(v, r, a, parsed) for a, v in r["values"].items()]
            vals = [v for v in vals if v is not None]
            if not vals:
                m = re.search(r"R?\$?\s*([\d.]+,\d{2})", " ".join(r["label_texts"]))
                vals = [parse_brl(m.group(1))] if m else []
            parsed.total_row_amounts += vals
            if len(vals) > 1:
                parsed.stated_total = float(sum(vals))
                parsed.notes.append("grand total row printed per agency; stated total is its sum")
            elif vals:
                parsed.stated_total = vals[0]
            elif any(r.get("uncached")):
                parsed.notes.append("grand total is a formula without a stored value")
            continue
        if r["kind"] == "agency_total":
            for a, v in r["values"].items():
                v = total_cell_amount(v, r, a, parsed)
                if v is not None:
                    parsed.agency_totals[a] = v
                    parsed.total_row_amounts.append(v)
            continue
        for a, raw in r["values"].items():
            if isinstance(raw, str) and not re.search(r"\d", raw):
                if raw.strip(" -–R$"):
                    parsed.notes.append(f"text in an amount cell: {clean(raw)!r}")
                continue
            if isinstance(raw, str) and INT_AMOUNT_RE.fullmatch(clean(raw)):
                whole.append((r, a, clean(raw), len(parsed.lines)))
                continue
            if isinstance(raw, str) and US_AMOUNT_RE.fullmatch(clean(raw)):
                parsed.notes.append(f"US-format amount outside a total row not read: "
                                    f"{clean(raw)!r} at {cell_position(r, a)}")
                continue
            v = parse_brl(raw)
            if v is None or v == 0:
                continue
            parsed.lines.append(dict(agency=a, section=r["section"], category=r["group"],
                                     subcategory=r["item"], amount_brl=v,
                                     page_or_sheet=r["where"]))
    if whole:
        resolve_whole_numbers(whole, parsed)


# ---------------------------------------------------------------------------
# PDF: ruled tables
# ---------------------------------------------------------------------------
def rotated_text(page, bbox) -> tuple[str, bool]:
    """Text of a table cell written vertically, read bottom-to-top with spaces at wide gaps.
    Covers rotated glyphs and upright letters stacked one per line. ('', False) otherwise."""
    x0, top, x1, bottom = bbox
    chars = [c for c in page.chars if c["text"].strip()
             and x0 <= (c["x0"] + c["x1"]) / 2 <= x1 and top <= (c["top"] + c["bottom"]) / 2 <= bottom]
    if len(chars) < 2:
        return "", False
    rot = [c for c in chars if not c.get("upright", True)]
    if len(rot) >= 0.5 * len(chars):
        seq = rot
    else:
        width = max(c["x1"] - c["x0"] for c in chars)
        tops = sorted(round(c["top"]) for c in chars)
        stacked = (len(chars) >= 3 and max(c["x0"] for c in chars) - min(c["x0"] for c in chars)
                   < 1.2 * width and len(set(tops)) >= len(chars) - 1)
        if not stacked:
            return "", False
        seq = chars
    seq = sorted(seq, key=lambda c: -(c["top"] + c["bottom"]) / 2)
    gaps = [prev["top"] - cur["bottom"] for prev, cur in zip(seq, seq[1:])]
    base = float(np.median(gaps)) if gaps else 0.0
    out = seq[0]["text"]
    for gap, cur in zip(gaps, seq[1:]):
        out += (" " if gap > base + 0.35 * max(cur.get("size", 8), 1) else "") + cur["text"]
    return clean(out), True


def page_is_supplier_list(text: str) -> bool:
    head = fold(text[:200])
    return "fornecedores" in head and "cnpj" in head


def _is_amount_text(c: str | None) -> bool:
    return bool(c) and bool(MONEY_RE.match(re.sub(r"[\s​R$]+", "", c) or "x"))


def parse_pdf_ruled(pdf, parsed: Parsed) -> bool:
    """True when the document has a ruled summary table with an agency header.

    Geometry, not column indices, links the pieces: amount cells go to the agency whose header
    cell is horizontally nearest (some 2019 tables put amounts in columns offset from the header
    cells), and label cells go to the label column of the header table whose left edge is
    nearest (a continuation table on a later page can have a different column grid)."""
    raw_rows: list[dict] = []
    centers: list[tuple[float, str]] | None = None
    canon_lefts: list[float] = []
    carry: list[str] = []
    done = False
    for pno, page in enumerate(pdf.pages):
        if done:
            break
        text = page.extract_text() or ""
        if page_is_supplier_list(text):
            if centers:
                break
            continue
        for table in page.find_tables():
            grid = table.extract()
            if not grid:
                continue
            ncols = max(len(r) for r in grid)
            lefts = []
            for j in range(ncols):
                xs = [row.cells[j][0] for row in table.rows
                      if j < len(row.cells) and row.cells[j] is not None]
                lefts.append(min(xs) if xs else float("inf"))
            header = next((i for i, row in enumerate(grid)
                           if sum(agency_name(c) is not None for c in row) >= 2), None)
            continuation = bool(centers)
            if header is not None:
                row, boxes = grid[header], table.rows[header].cells
                heads = [j for j, c in enumerate(row) if agency_name(c)]
                centers = [((boxes[j][0] + boxes[j][2]) / 2, agency_name(row[j])) for j in heads]
                agency_left = min(boxes[j][0] for j in heads)
                start = header + 1
                if parsed.title_period is None:
                    parsed.title_period = title_period(text) or title_period(clean(row[0]))
            elif centers and any(_is_amount_text(c) or clean(c) for r in grid for c in r):
                start = 0
            else:
                continue
            if pno + 1 not in parsed.table_pages:
                parsed.table_pages.append(pno + 1)
            amount_cols = [j for i in range(start, len(grid)) for j, c in enumerate(grid[i])
                           if _is_amount_text(c)]
            n_label = min([j for j in range(ncols) if lefts[j] >= agency_left - 2] + amount_cols
                          + [ncols])
            if not canon_lefts:
                canon_lefts = lefts[:n_label]
                carry = [""] * n_label
            slot_of = {j: min(range(len(canon_lefts)), key=lambda k: abs(canon_lefts[k] - lefts[j]))
                       for j in range(n_label)}
            for i in range(start, len(grid)):
                raw, boxes = grid[i], table.rows[i].cells
                cells: list[dict | None] = [None] * len(canon_lefts)
                covered = [False] * len(canon_lefts)
                for j in range(n_label):
                    if raw[j] is None:
                        # a merged cell from the left spans this column: blank, not carried
                        if any(boxes[k] is not None and boxes[k][2] > lefts[j] + 1
                               for k in range(j)):
                            covered[slot_of[j]] = True
                        continue
                    t, rot = rotated_text(page, boxes[j]) if boxes[j] else ("", False)
                    t = t if rot else clean(raw[j])
                    k = slot_of[j]
                    if cells[k] is not None and cells[k]["text"]:
                        t = f"{cells[k]['text']} {t}".strip()
                    cells[k] = {"text": t, "rotated": rot}
                labels = []
                for k, c in enumerate(cells):
                    if c is None:
                        labels.append("" if covered[k] else carry[k])
                    elif c["text"] == "" and continuation and i == start:
                        labels.append(carry[k])
                        cells[k] = None
                    else:
                        labels.append(c["text"])
                carry = labels
                values: dict[str, str] = {}
                pos: dict[str, str] = {}
                for j in range(n_label, len(raw)):
                    if raw[j] is None or not clean(raw[j]) or not boxes[j]:
                        continue
                    xc = (boxes[j][0] + boxes[j][2]) / 2
                    agency = min(centers, key=lambda c: abs(c[0] - xc))[1]
                    if agency in values:
                        raise ValueError(f"p{pno + 1}: two cells for {agency} in row {raw}")
                    values[agency] = raw[j]
                    pos[agency] = f"x {boxes[j][0]:.0f}, y {boxes[j][1]:.0f}"
                own = [c["text"] for c in cells if c and c["text"]]
                raw_rows.append(dict(kind=row_kind(own), own_label=bool(own), cells=cells,
                                     labels=labels, values=values, where=f"p{pno + 1}",
                                     label_texts=own, uncached=[], pos=pos))
                if raw_rows[-1]["kind"] == "grand_total":
                    done = True
                    break
            if done:
                break
    if not centers:
        return False
    parsed.agencies = [a for _, a in sorted(centers)]

    # Rows without any label of their own that repeat the running sums are total rows printed
    # without a label (February 2019).
    # A US-format cell is left to lines_from_rows, which reads it on a total row only.
    running = {a: 0.0 for a in parsed.agencies}
    for r in raw_rows:
        vals = {a: parse_brl(v) for a, v in r["values"].items()
                if re.search(r"\d", v) and not US_AMOUNT_RE.fullmatch(clean(v))}
        vals = {a: v for a, v in vals.items() if v}
        if r["kind"] == "data" and not r["own_label"] and vals:
            if len(vals) >= 2 and all(within_tol(v, running[a]) for a, v in vals.items()):
                r["kind"] = "agency_total"
                parsed.notes.append("unlabeled per-agency total row")
                continue
            if len(vals) == 1 and within_tol(next(iter(vals.values())), sum(running.values())):
                r["kind"] = "grand_total"
                parsed.notes.append("unlabeled grand total row")
                continue
        if r["kind"] == "data":
            for a, v in vals.items():
                running[a] += v

    width = len(canon_lefts)
    used = [k for k in range(width)
            if any(r["cells"][k] and r["cells"][k]["text"] for r in raw_rows if r["kind"] == "data")]
    item_c = used[-1] if used else None
    group_c = used[-2] if len(used) >= 2 else None
    section_cs = used[:-2]
    for r in raw_rows:
        r["group"] = r["labels"][group_c] if group_c is not None else ""
        r["item"] = r["labels"][item_c] if item_c is not None else ""
        r["sec_raw"] = [r["cells"][k] for k in section_cs]
    resolve_sections(raw_rows, len(section_cs))
    lines_from_rows(raw_rows, parsed.agencies, parsed)
    return True


# ---------------------------------------------------------------------------
# PDF: unruled layout (2014-2018)
# ---------------------------------------------------------------------------
def _words(page) -> list[dict]:
    upright = page.filter(lambda o: o.get("object_type") != "char" or o.get("upright", True))
    words = upright.extract_words(x_tolerance=1.5, y_tolerance=2)
    merged: list[dict] = []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if merged:
            p = merged[-1]
            numeric = re.fullmatch(r"[\d.,]+", w["text"]) and re.fullmatch(r"[\d.,]+", p["text"])
            if numeric and abs(p["top"] - w["top"]) < 2 and 0 <= w["x0"] - p["x1"] < 2.5:
                p.update(text=p["text"] + w["text"], x1=w["x1"])
                continue
        merged.append(dict(w))
    return merged


def _cluster_lines(words: list[dict], tol: float = 2.5) -> list[list[dict]]:
    lines: list[list[dict]] = []
    for w in sorted(words, key=lambda w: w["top"]):
        if lines and abs(lines[-1][0]["top"] - w["top"]) <= tol:
            lines[-1].append(w)
        else:
            lines.append([w])
    return [sorted(line, key=lambda w: w["x0"]) for line in lines]


def _isolated_word(w: dict, words: list[dict], gap: float = 4.0) -> bool:
    """No other word on w's line within `gap` points of either side. A CNPJ, a date, a phone
    number or a percentage can carry thousands-style dots, but it comes glued to a '/', a '-', a
    '%' or another group of digits ("11.222.333/0001-44", "12.500 %"); a cell of the amount
    columns is set apart from its neighbours by the column spacing. A word that passes is only a
    whole-number candidate: resolve_whole_numbers reads it only when the printed totals need it,
    which covers what spacing cannot (a lone CNPJ root, a '%' printed a column away)."""
    for o in words:
        if o is w or abs(o["top"] - w["top"]) >= 2:
            continue
        if 0 <= w["x0"] - o["x1"] < gap or 0 <= o["x0"] - w["x1"] < gap:
            return False
    return True


def parse_pdf_unruled(pdf, parsed: Parsed) -> bool:
    cols: list[tuple[float, str]] | None = None
    first_x0 = 0.0
    label_lines: list[dict] = []
    done = False
    for pno, page in enumerate(pdf.pages):
        if done:
            break
        text = page.extract_text() or ""
        if page_is_supplier_list(text):
            if cols:
                break
            continue
        words = _words(page)
        lines = _cluster_lines(words)
        top_limit = 0.0
        for line in lines:
            ags = [w for w in line if agency_name(w["text"])]
            if len(ags) >= 2:
                cols = [((w["x0"] + w["x1"]) / 2, agency_name(w["text"])) for w in ags]
                first_x0 = min(w["x0"] for w in ags)
                top_limit = max(w["bottom"] for w in ags)
                if parsed.title_period is None:
                    parsed.title_period = title_period(text)
                break
        if not cols:
            continue
        body = [w for w in words if w["top"] > top_limit]
        money, labels = [], []
        for w in body:
            t = w["text"]
            if MONEY_RE.match(t):
                money.append(w)
            elif w["x0"] >= first_x0 - 2:
                # The amount columns. A word here is either an amount or dropped, so a dropped
                # word that carries digits is recorded: May 2017's "14.250" was dropped silently.
                # A whole-number cell and a US-format amount go on to lines_from_rows, which
                # decides whether they are read (the printed totals; a total row).
                if (INT_AMOUNT_RE.match(t) and _isolated_word(w, body)) or US_AMOUNT_RE.match(t):
                    money.append(w)
                elif re.search(r"\d", t):
                    parsed.notes.append(f"p{pno + 1}: word with digits in the amount columns "
                                        f"not read as an amount: {t!r}")
            elif t != "R$" and not re.fullmatch(r"\d{1,2}", t):
                labels.append(w)
        page_lines = []
        for line in _cluster_lines(labels):
            joined = " ".join(w["text"] for w in line)
            if re.search(r"custos|despesas com publicidade|caixa", fold(joined)):
                continue
            segs = [[line[0]]]
            for w in line[1:]:
                if w["x0"] - segs[-1][-1]["x1"] > 12:
                    segs.append([w])
                else:
                    segs[-1].append(w)
            page_lines.append(dict(
                top=min(w["top"] for w in line), bottom=max(w["bottom"] for w in line),
                x0=line[0]["x0"], segs=[(s[0]["x0"], " ".join(w["text"] for w in s)) for s in segs],
                values={}, pos={}, where=f"p{pno + 1}", dist=0.0))
        if not page_lines:
            continue
        for w in money:
            yc = (w["top"] + w["bottom"]) / 2
            best = min(page_lines, key=lambda ln: abs((ln["top"] + ln["bottom"]) / 2 - yc))
            d = abs((best["top"] + best["bottom"]) / 2 - yc)
            xc = (w["x0"] + w["x1"]) / 2
            agency = min(cols, key=lambda c: abs(c[0] - xc))[1]
            if agency in best["values"]:
                best["values"][agency] += " + " + w["text"]
            else:
                best["values"][agency] = w["text"]
                best["pos"][agency] = f"x {w['x0']:.0f}, y {w['top']:.0f}"
            best["dist"] = max(best["dist"], d)
        label_lines.extend(page_lines)
        parsed.table_pages.append(pno + 1)
        if any(row_kind([s for _, s in ln["segs"]]) == "grand_total" for ln in page_lines):
            done = True
    if not cols:
        return False
    parsed.agencies = [a for _, a in sorted(cols)]

    group_x = min(ln["x0"] for ln in label_lines)
    rows, group, prev = [], "", None
    for ln in label_lines:
        g = " ".join(t for x, t in ln["segs"] if x < group_x + 60)
        it = " ".join(t for x, t in ln["segs"] if x >= group_x + 60)
        kind = row_kind([g, it])
        if kind == "data" and not ln["values"] and not g and prev is not None \
                and prev["kind"] == "data" and prev["item"] and prev["where"] == ln["where"] \
                and ln["top"] - prev["bottom"] < 6:
            prev["item"] = f"{prev['item']} {it}"
            continue
        if kind == "data" and g:
            group = g
        if any(" + " in v for v in ln["values"].values()):
            raise ValueError(f"two amounts for one agency on the line {g} {it}: {ln['values']}")
        row = dict(kind=kind, section="", group=group if kind == "data" else "", item=it,
                   values=ln["values"], where=ln["where"], label_texts=[g, it], uncached=[],
                   bottom=ln["bottom"], pos=ln["pos"])
        if kind == "data" and ln["values"] and ln["dist"] > AMBIGUOUS_PT:
            parsed.ambiguous_rows += 1
        rows.append(row)
        prev = row
        if kind == "grand_total":
            break
    lines_from_rows(rows, parsed.agencies, parsed)
    return True


def parse_pdf(b: bytes, member: str, method: str) -> Parsed:
    parsed = Parsed(method=method, member=member)
    with pdfplumber.open(io.BytesIO(b)) as pdf:
        full = " ".join((pg.extract_text() or "") for pg in pdf.pages[:6])
        parsed.no_payments_stated = "nao existem pagamentos" in fold(full)
        if parse_pdf_ruled(pdf, parsed):
            return parsed
        parsed = Parsed(method=method, member=member,
                        no_payments_stated=parsed.no_payments_stated)
        if parse_pdf_unruled(pdf, parsed):
            return parsed
        first = pdf.pages[0]
        if len(first.chars) < 20 and first.images:
            parsed.status = "unreadable_scan"
            return parsed
    raise ValueError(f"{member}: no summary table with an agency header found")


# ---------------------------------------------------------------------------
# Spreadsheets
# ---------------------------------------------------------------------------
def _grid_openpyxl(b: bytes) -> list[tuple[str, list[list], list[list[bool]], list]]:
    vals = openpyxl.load_workbook(io.BytesIO(b), data_only=True)
    forms = openpyxl.load_workbook(io.BytesIO(b), data_only=False)
    out = []
    for ws, wf in zip(vals.worksheets, forms.worksheets):
        grid = [[c.value for c in row] for row in ws.iter_rows()]
        uncached = [[(c.value is None and isinstance(f.value, str) and f.value.startswith("="))
                     for c, f in zip(row, frow)]
                    for row, frow in zip(ws.iter_rows(), wf.iter_rows())]
        merges = [(m.min_row - 1, m.max_row, m.min_col - 1, m.max_col)
                  for m in ws.merged_cells.ranges]
        out.append((ws.title, grid, uncached, merges))
    return out


def _grid_xlrd(b: bytes) -> list[tuple[str, list[list], list[list[bool]], list]]:
    book = xlrd.open_workbook(file_contents=b, formatting_info=True)
    out = []
    for sh in book.sheets():
        grid = [[(None if sh.cell_type(r, c) in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK)
                  else sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)]
        uncached = [[False] * sh.ncols for _ in range(sh.nrows)]
        out.append((sh.name, grid, uncached, list(sh.merged_cells)))
    return out


def parse_spreadsheet(b: bytes, member: str, kind: str, method: str) -> Parsed:
    sheets = _grid_xlrd(b) if kind == "xls" else _grid_openpyxl(b)
    for name, grid, uncached, merges in sheets:
        header = next((i for i, row in enumerate(grid)
                       if sum(agency_name(v) is not None for v in row) >= 2), None)
        if header is None:
            continue
        parsed = Parsed(method=method, member=member)
        agency_cols = [(j, agency_name(v)) for j, v in enumerate(grid[header]) if agency_name(v)]
        parsed.agencies = [a for _, a in agency_cols]
        first_agency = agency_cols[0][0]
        for row in grid[:header + 1]:
            for v in row:
                if isinstance(v, str) and parsed.title_period is None:
                    parsed.title_period = title_period(v)

        owner: dict[tuple[int, int], tuple[int, int]] = {}
        for r0, r1, c0, c1 in merges:
            for r in range(r0, r1):
                for c in range(c0, c1):
                    owner[(r, c)] = (r0, c0)

        def row_labels(r: int) -> list[str]:
            """Label texts of row r; a cell merged with a cell to its left is blank."""
            out, seen = [], set()
            for c in range(first_agency):
                rr, cc = owner.get((r, c), (r, c))
                v = grid[rr][cc] if cc < len(grid[rr]) else None
                out.append(clean(v) if isinstance(v, str) and (rr, cc) not in seen else "")
                seen.add((rr, cc))
            return out

        body = []
        for r in range(header + 1, len(grid)):
            labels = row_labels(r)
            kind_r = row_kind(labels)
            body.append((r, labels, kind_r))
            if kind_r == "grand_total":
                break
        used = [c for c in range(first_agency)
                if any(lab[c] for _, lab, k in body if k == "data")]
        if not used:
            raise ValueError(f"{member}/{name}: no label columns")
        item_c = used[-1]
        group_c = used[-2] if len(used) >= 2 else None
        section_cs = used[:-2]
        label_merges = any(c0 < first_agency and r1 - r0 > 1 for r0, r1, c0, c1 in merges)

        rows, group, sections = [], "", [""] * len(section_cs)
        for r, labels, kind_r in body:
            values, flags = {}, []
            for c, a in agency_cols:
                v = grid[r][c] if c < len(grid[r]) else None
                if v is None and c < len(uncached[r]) and uncached[r][c]:
                    flags.append(a)
                values[a] = v
            if kind_r == "data":
                g = labels[group_c] if group_c is not None else ""
                it = labels[item_c]
                if not label_merges:
                    if g:
                        group = g
                    elif it:
                        g = group
                    for k, c in enumerate(section_cs):
                        sections[k] = labels[c] or sections[k]
                    secs = sections
                else:
                    secs = [labels[c] for c in section_cs]
                if not (g or it) and all(v in (None, "") for v in values.values()):
                    continue
                if flags:
                    raise ValueError(f"{member}/{name} row {r + 1}: formula without a stored "
                                     f"value in an amount cell for {flags}")
            else:
                g, it, secs = "", "", []
            rows.append(dict(kind=kind_r, section=" > ".join(s for s in secs if s), group=g,
                             item=it, values=values, where=name, label_texts=labels,
                             uncached=flags, pos={a: f"cell {get_column_letter(c + 1)}{r + 1}"
                                                  for c, a in agency_cols}))
        lines_from_rows(rows, parsed.agencies, parsed)
        return parsed
    raise ValueError(f"{member}: no sheet with an agency header")


# ---------------------------------------------------------------------------
# One month
# ---------------------------------------------------------------------------
def parse_slot(path: Path, slot: str) -> tuple[list[Parsed], list[str]]:
    """Every parseable document in a downloaded file, and notes on what could not be used."""
    out, notes = [], []
    for member, b in unpack(path):
        kind = kind_of(member, b)
        if "fornecedor" in fold(member) and path.suffix.lower() == ".zip":
            continue
        if kind == "pdf":
            method = "pdf_in_xlsx_slot" if slot == "editable" else "text_pdf"
            out.append(parse_pdf(b, member, method))
        elif kind in ("xlsx", "template", "xls"):
            out.append(parse_spreadsheet(b, member, kind, "template" if kind == "template"
                                         else "xlsx"))
        elif kind == "html":
            notes.append(f"{slot} link returns a web page, not a document")
        else:
            notes.append(f"{slot} member {member} of unknown type")
    return out, notes


def within_tol(a: float, b: float) -> bool:
    """Whether a row printed without a label repeats the running sums, which makes it a total
    row (parse_pdf_ruled). This only classifies the row: its amounts then become printed totals
    and are checked in cents like any other, so the looser match accepts no disagreement."""
    return abs(a - b) <= max(TOL_ABS, TOL_REL * abs(b))


VALIDATED_EXACT = "validated_exact"
VALIDATED_WITHIN = f"validated_within_{MAX_GAP_CENTS}_centavos"
NOT_VALIDATED_DOCUMENT = "not_validated_document"
NOT_VALIDATED_PARSER = "not_validated_parser"
NO_STATED_TOTAL = "no_stated_total"
VALIDATED_STATUSES = (VALIDATED_EXACT, VALIDATED_WITHIN)


def gap_fields(checks: list[tuple[str, int]], has_total: bool,
               detail_only: list[tuple[str, int]] = ()) -> dict:
    """validation_status and the gap columns from a month's (check, gap in cents) pairs.

    A failing month starts as not_validated_parser; build_month moves it to
    not_validated_document only on the evidence of explain_failure. detail_only pairs are
    reported in gap_detail without entering the status."""
    worst = max((abs(g) for _, g in checks), default=None)
    allowed = [(c, g) for c, g in checks if 0 < abs(g) <= MAX_GAP_CENTS]
    if not has_total:
        status = NO_STATED_TOTAL
    elif worst > MAX_GAP_CENTS:
        status = NOT_VALIDATED_PARSER
    else:
        status = VALIDATED_WITHIN if worst else VALIDATED_EXACT
    return dict(validation_status=status, validated=status in VALIDATED_STATUSES,
                max_gap_cents=pd.NA if worst is None else worst,
                # The same meaning as in scrape_caixa_scan_advertising.py: the month was accepted
                # only through the allowance. tolerance_detail still lists every check that used
                # it, so a failing month shows those too.
                tolerance_flagged=status == VALIDATED_WITHIN,
                tolerance_detail="; ".join(f"{c} {g:+d}" for c, g in allowed),
                gap_detail="; ".join(f"{c} {g:+d}" for c, g in list(checks) + list(detail_only)
                                     if g))


# ---------------------------------------------------------------------------
# Whose gap is it: the image check
# ---------------------------------------------------------------------------
IMAGE_DPI = 200
# OCR can put several cells of a row in one box, spaced ("14.191.628,55 11.150.869,67") or run
# together ("6.162.385,8312.850.749,717.321.259,75", October 2017's TOTAL row). A run of digits,
# dots and commas that is one amount is read as it stands; any other run with a decimal comma is
# cut into amounts of two decimals, as the cells are normally printed, and is kept only if those
# amounts rebuild it character for character. OCR's known misreads of one amount are read as
# that amount: the decimal comma read as a dot ("26.453.70"), a space before the cents
# ("26.453, 70", joined first), US format ("26,453.70"); a thousands-grouped integer ("14.250")
# is read as the whole number the page shows. Any other run with a decimal comma, or of 4 or more
# digits with a dot or comma, is recorded (image_unreadable), because an amount could hide in it
# and the evidence is then incomplete.
OCR_RUN_RE = re.compile(r"[\d.,]+")
OCR_AMOUNT_RE = re.compile(r"\d{1,3}(?:\.\d{3})*,\d{1,4}")
OCR_CELL_RE = re.compile(r"\d{1,3}(?:\.\d{3})*,\d{2}")
OCR_DOT_CENTS_RE = re.compile(r"\d{1,3}(?:\.\d{3})+\.\d{2}")
OCR_SPACED_CENTS_RE = re.compile(r"(?<![\d.,])(\d{1,3}(?:\.\d{3})*),\s+(\d{2})(?![\d.,])")
# A word of an OCR box, keeping "NOVA / SB" whole when OCR spaces the slash.
OCR_WORD_RE = re.compile(r"\S+?\s*/\s*\S+|\S+")
# A Brazilian-format amount in the text layer where the image is not read (text_layer_outside).
TEXT_AMOUNT_RE = re.compile(r"(?<![\d.,])\d{1,3}(?:\.\d{3})*,\d{2}(?![\d,])")
_OCR_ENGINE = None


def ocr_engine():
    """RapidOCR, loaded on first use: only a month that fails validation is read as an image."""
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        from rapidocr_onnxruntime import RapidOCR
        _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


def member_bytes(path: Path, member: str) -> bytes | None:
    return next((b for m, b in unpack(path) if m == member), None)


def ocr_pages(pdf_bytes: bytes, table_pages: list[int]) -> tuple[list[dict], list[int]]:
    """OCR boxes of the pages that hold the summary table, rendered with PyMuPDF at IMAGE_DPI one
    page at a time, and the page numbers read. Those are the pages the parser read the table from
    and every page without a text layer, which the parser cannot have read at all. No other page
    is imaged: the supplier lists the parser skips, and any page after the one that carries the
    TOTAL GERAL, where the parser stops and never looks. At about 25 seconds a page, July 2016's
    20 supplier pages would cost minutes, so those pages are covered by a scan of their text
    layer instead (text_layer_outside), as is the region above the header row of an imaged page,
    which image_amounts does not count as table."""
    import fitz
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        pages = [i for i, pg in enumerate(pdf.pages)
                 if i + 1 in table_pages or not pg.chars]
    engine = ocr_engine()
    boxes: list[dict] = []
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    try:
        for i in pages:
            pix = doc[i].get_pixmap(dpi=IMAGE_DPI)
            img = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, pix.n)
            img = np.ascontiguousarray(img[:, :, 2::-1])       # RGB(A) to the BGR RapidOCR reads
            del pix
            result, _ = engine(img)
            del img
            for quad, text, _conf in result or []:
                xs, ys = [p[0] for p in quad], [p[1] for p in quad]
                boxes.append(dict(text=str(text), x0=min(xs), x1=max(xs), top=min(ys),
                                  bottom=max(ys), page=i + 1))
    finally:
        doc.close()
    return boxes, [i + 1 for i in pages]


def _header_top(page) -> float | None:
    """Top of the first text-layer line with two agency names, or None."""
    for line in _cluster_lines(_words(page)):
        if sum(agency_name(w["text"]) is not None for w in line) >= 2:
            return min(w["top"] for w in line)
    return None


def text_layer_outside(pdf_bytes: bytes, imaged: list[int]) -> list[str]:
    """Nonzero Brazilian-format amounts in the text layer where the image is not read: every page
    not imaged (ocr_pages), and on an imaged page the region above its header row. The parser
    does not read these places either, so an amount there would be invisible to both; reading a
    text layer costs little next to OCR. A page without a text layer is imaged whole."""
    hits = []
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for i, page in enumerate(pdf.pages):
            if not page.chars:
                continue
            where, words = f"p{i + 1}", page.extract_words(x_tolerance=1.5, y_tolerance=2)
            if i + 1 in imaged:
                top = _header_top(page)
                if top is None:
                    continue
                where += ", text layer above the header row"
                words = [w for w in words if w["bottom"] <= top]
            for w in words:
                for m in TEXT_AMOUNT_RE.finditer(w["text"]):
                    if cents([parse_brl(m.group(0))]):
                        hits.append(f"{m.group(0)} ({where})")
    return hits


def _box_x(box: dict, start: int, end: int) -> float:
    """Horizontal centre of characters start:end of an OCR box, placed by their offsets within
    the box text, for a box that holds several cells."""
    return box["x0"] + (box["x1"] - box["x0"]) * (start + end) / 2 / max(len(box["text"]), 1)


def _box_tokens(box: dict, pattern: re.Pattern) -> list[tuple[str, float]]:
    """(match, horizontal centre) for each match of pattern in an OCR box."""
    return [(m.group(0), _box_x(box, m.start(), m.end())) for m in pattern.finditer(box["text"])]


def _box_amounts(box: dict) -> tuple[list[tuple[str, float]], list[str]]:
    """(amount in Brazilian format, horizontal centre) for each amount in an OCR box, and the runs
    that could not be read as amounts (see OCR_RUN_RE)."""
    joined = dict(box, text=OCR_SPACED_CENTS_RE.sub(r"\1,\2", box["text"]))
    found, unreadable = [], []
    for run in OCR_RUN_RE.finditer(joined["text"]):
        text = run.group(0).strip(".,")
        if not text:
            continue
        start = run.start() + run.group(0).index(text)
        centre = _box_x(joined, start, start + len(text))
        if OCR_AMOUNT_RE.fullmatch(text) or INT_AMOUNT_RE.match(text):
            found.append((text, centre))
        elif OCR_DOT_CENTS_RE.fullmatch(text):
            found.append((f"{text[:-3]},{text[-2:]}", centre))
        elif US_AMOUNT_RE.match(text):
            found.append((text.replace(",", "_").replace(".", ",").replace("_", "."), centre))
        elif "," in text and "".join(c.group(0) for c in OCR_CELL_RE.finditer(text)) == text:
            found += [(c.group(0), _box_x(joined, start + c.start(), start + c.end()))
                      for c in OCR_CELL_RE.finditer(text)]
        elif "," in text or ("." in text and sum(ch.isdigit() for ch in text) >= 4):
            unreadable.append(text)
    return found, unreadable


def _total_label(text: str) -> bool:
    """An OCR box that labels a total row, by row_kind's words; words with digits and "R$" are
    left out, so a box that runs the label into its amounts still counts."""
    words = [w for w in text.split() if not re.search(r"\d", w) and w != "R$"]
    return row_kind([" ".join(words)]) != "data"


def image_amounts(boxes: list[dict]) -> tuple[list[tuple[str | None, int, int, bool]], list[str],
                                               list[str]]:
    """Every nonzero amount OCR read below the header row, as (nearest agency header, cents, page,
    on a total row); the runs it could not read as amounts; and the amounts it read above the
    header row.

    The headers are the first OCR row with two agency names, on that page or the last page that
    had one; on a page with a header row what lies above it is the title, not table, and an amount
    there is returned apart (incomplete evidence, like the text layer outside the image). An
    amount is on a total row when a box labelled TOTAL (TOTAL GERAL, TOTAL POR ...) sits at its
    height. The nearest header only says where an amount sits, for the log: it is not a column
    reading. The ruled tables right-align amounts in cells whose header is centred, so a centre
    rule puts July 2023's amounts one column to the right of the spreadsheet's, and in the
    unruled layout it is the parser's own rule, so agreeing with it would prove nothing."""
    out: list[tuple[str | None, int, int, bool]] = []
    unreadable: list[str] = []
    above: list[str] = []
    centers: list[tuple[float, str]] = []
    for page in sorted({b["page"] for b in boxes}):
        on_page = sorted((b for b in boxes if b["page"] == page), key=lambda b: b["top"])
        heads = [(b, xc, agency_name(tok)) for b in on_page
                 for tok, xc in _box_tokens(b, OCR_WORD_RE) if agency_name(tok)]
        header_top = None
        for b, _, _ in heads:
            row = [(xc, a) for bb, xc, a in heads if abs(bb["top"] - b["top"]) < 0.6 * (
                b["bottom"] - b["top"])]
            if len(row) >= 2:
                centers, header_top = sorted(row), b["top"]
                break
        labels = [b for b in on_page if _total_label(b["text"])]
        for b in on_page:
            found, bad = _box_amounts(b)
            if header_top is not None and b["top"] <= header_top:
                above += [f"{tok} (p{page}, OCR above the header row)" for tok, _ in found
                          if cents([parse_brl(tok)])]
                continue
            unreadable += [f"{t!r} (p{page})" for t in bad]
            yc = (b["top"] + b["bottom"]) / 2
            on_total = any(abs((t["top"] + t["bottom"]) / 2 - yc) < 0.5 * (t["bottom"] - t["top"])
                           for t in labels)
            for tok, xc in found:
                c = cents([parse_brl(tok)])
                if c:
                    agency = min(centers, key=lambda h: abs(h[0] - xc))[1] if centers else None
                    out.append((agency, c, page, on_total))
    return out, unreadable, above


def brl_text(c: int) -> str:
    return f"{c / 100:,.2f}"


def image_evidence(pdf_bytes: bytes, image_doc: Parsed, lines: list[dict]) -> tuple[dict,
                                                                                     list[int]]:
    """Compare the multiset of amounts the page images show with the amounts the parser used.
    Returns the image_* columns and the cents of every amount printed on a total row (those the
    image shows on a row labelled TOTAL, and any printed total of the parser's that is not among
    them), for the refusal in explain_failure.

    Each printed total of the imaged document is paired with an image amount of the same cents,
    one on a total row if there is one; each parser line then with one of the rest, one off a
    total row if there is one. A line that can only take an amount the image shows on a total row
    is a total read as data (image_lines_on_total_rows). An image amount still unpaired is one the
    parser did not use (image_unparsed). A parser line or printed total left unpaired is one OCR
    did not find (image_missing, image_totals_missing), which makes the image reading incomplete,
    not the parser wrong. The parser's zero cells are not lines and the image's zeros are dropped,
    so neither is compared."""
    boxes, pages = ocr_pages(pdf_bytes, image_doc.table_pages)
    found, unreadable, above = image_amounts(boxes)
    outside = above + text_layer_outside(pdf_bytes, pages)
    pool = list(found)

    def take(c: int, on_total: bool) -> tuple | None:
        k = next((i for i, f in enumerate(pool) if f[1] == c and f[3] == on_total), None)
        if k is None:
            k = next((i for i, f in enumerate(pool) if f[1] == c), None)
        return None if k is None else pool.pop(k)

    printed = [c for c in (cents([v]) for v in image_doc.total_row_amounts) if c]
    totals_missing = [c for c in printed if take(c, True) is None]
    missing, on_total_rows = [], []
    for x in lines:
        c = cents([x["amount_brl"]])
        f = take(c, False)
        if f is None:
            missing.append((x["agency"], c))
        elif f[3]:
            on_total_rows.append((x["agency"], c, f[2]))
    total_row_cents = [f[1] for f in found if f[3]]
    unmatched = list(total_row_cents)
    for c in printed:
        if c in unmatched:
            unmatched.remove(c)
        else:
            total_row_cents.append(c)
    return dict(
        image_check="ran", image_pages=",".join(f"p{p}" for p in pages),
        image_amounts=len(found), image_unparsed=len(pool),
        image_unparsed_brl="; ".join(f"{brl_text(c)} (p{p}, nearest header "
                                     f"{a or 'none'})" for a, c, p, _ in pool),
        image_missing=len(missing),
        image_missing_brl="; ".join(f"{brl_text(c)} ({a})" for a, c in missing),
        image_totals_missing="; ".join(brl_text(c) for c in totals_missing),
        image_unreadable="; ".join(unreadable),
        image_outside_amounts="; ".join(outside),
        image_lines_on_total_rows="; ".join(f"{brl_text(c)} ({a}, p{p})"
                                            for a, c, p in on_total_rows)), total_row_cents


def total_row_subset(target: int, amounts: list[int]) -> list[int]:
    """Some of the amounts (at least one) whose sum is within MAX_GAP_CENTS of target, or [].
    A document has a few total-row amounts (one per agency and the TOTAL GERAL), and sums past
    the target are dropped as the search goes, so trying every combination is cheap."""
    reach: dict[int, list[int]] = {0: []}
    for a in amounts:
        if a <= 0:
            continue
        for s, used in list(reach.items()):
            if s + a <= target + MAX_GAP_CENTS and s + a not in reach:
                reach[s + a] = used + [a]
    return next((used for s, used in reach.items()
                 if used and abs(s - target) <= MAX_GAP_CENTS), [])


def printed_totals_conflict(doc: Parsed, checks: list[tuple[str, int]]) -> str:
    """The second, independent sign that a gap is the document's: its printed totals disagree
    with each other. Returns a description when the printed agency totals do not add up to the
    printed grand total AND that disagreement accounts for every failing check (the lines agree
    with one side of it; April 2017 prints agency totals that sum to 10.00 less than its TOTAL
    GERAL, and its lines match the TOTAL GERAL). Otherwise ''. It is read from the parser's own
    reading of the totals, so it counts only with complete image evidence (explain_failure)."""
    if not doc.agency_totals or doc.stated_total is None:
        return ""
    diff = cents(doc.agency_totals.values()) - cents([doc.stated_total])
    if abs(diff) <= MAX_GAP_CENTS:
        return ""
    failing = [(c, g) for c, g in checks if abs(g) > MAX_GAP_CENTS]
    if any(not c.startswith(("total", "agency_total ")) for c, _ in failing):
        return ""
    total_gap = next((g for c, g in checks if c == "total"), None)
    agency_gaps = [g for c, g in checks if c.startswith("agency_total ")]
    # Lines equal to the TOTAL GERAL leave only agency gaps of one sign that add up to the
    # disagreement; lines equal to the agency totals leave only the TOTAL GERAL gap.
    one_side = ((total_gap is not None and abs(total_gap) <= MAX_GAP_CENTS
                 and sum(abs(g) for g in agency_gaps) <= abs(diff) + MAX_GAP_CENTS)
                or all(abs(g) <= MAX_GAP_CENTS for g in agency_gaps))
    if not one_side:
        return ""
    return (f"printed agency totals sum to {brl_text(cents(doc.agency_totals.values()))}, "
            f"printed TOTAL GERAL {brl_text(cents([doc.stated_total]))} ({diff:+d} cents), "
            f"which accounts for every failing check")


def _gaps_text(gaps: dict[str, int]) -> str:
    return "; ".join(f"{c} {g:+d}" for c, g in gaps.items())


def explain_failure(period: str, checks: list[tuple[str, int]], lines: list[dict],
                    totals_doc: Parsed, image_doc: tuple[bytes, Parsed] | None) -> dict:
    """not_validated_document or not_validated_parser for a failing month, with the evidence.
    The steps are those of the module docstring ("Whose gap is it").

    The page image decides what the parser missed: an amount the image shows and the parser did
    not use is a parser failure (1). The image can only clear the parser when its evidence is
    complete (2): OCR read amounts, found every parser amount and every printed total, left nothing
    unreadable, and no amount lies where the image is not read. Printed totals that disagree with
    each other are no substitute for that evidence, since the parser read those totals itself and a
    misread total produces the same disagreement; a month whose image could not be read, or read
    incompletely, is a parser failure whatever its totals say. A gap that equals printed total-row
    amounts, or a line the image shows only on a TOTAL row, is refused (3). With complete
    evidence an unclosed TOTAL GERAL, or for an agency-only failure a printed-totals conflict,
    makes the month a document candidate (4), and REVIEWED_DOCUMENT_MONTHS decides (5): listed
    with exactly these failing gaps, not_validated_document; otherwise not_validated_parser,
    "unreviewed document candidate"."""
    conflict = printed_totals_conflict(totals_doc, checks)
    out = dict(printed_totals_conflict=conflict)

    def parser(reason: str) -> dict:
        return out | dict(validation_status=NOT_VALIDATED_PARSER, validation_reason=reason)

    if image_doc is None:
        out.update(image_check="no PDF to render")
        return parser("image check could not run: no PDF of this month parsed")
    try:
        ev, total_row_cents = image_evidence(image_doc[0], image_doc[1], lines)
    except Exception as e:                                   # noqa: BLE001 - recorded, not hidden
        out.update(image_check=f"failed: {type(e).__name__}: {e}")
        return parser(f"image check could not run: {type(e).__name__}: {e}")
    out.update(ev)
    if ev["image_unparsed"]:                                                            # step 1
        return parser(f"the image shows {ev['image_unparsed']} amount(s) the parser did not use: "
                      f"{ev['image_unparsed_brl']}")
    incomplete = [x for x in (                                                          # step 2
        not ev["image_amounts"] and "OCR read no amount",
        ev["image_missing"] and f"OCR did not find {ev['image_missing']} parser amount(s): "
                                f"{ev['image_missing_brl']}",
        ev["image_totals_missing"] and f"printed totals not found in the image: "
                                       f"{ev['image_totals_missing']}",
        ev["image_unreadable"] and f"OCR text not readable as amounts: {ev['image_unreadable']}",
        ev["image_outside_amounts"] and f"amounts where the image is not read: "
                                        f"{ev['image_outside_amounts']}") if x]
    if incomplete:
        reason = "image evidence incomplete: " + "; ".join(incomplete)
        if conflict:
            reason += (f"; the printed totals disagree ({conflict}), which does not count "
                       f"without complete image evidence")
        return parser(reason)
    refused = []                                                                        # step 3
    if ev["image_lines_on_total_rows"]:
        refused.append(f"the parser used as lines amounts the image shows only on a TOTAL row: "
                       f"{ev['image_lines_on_total_rows']}")
    total_gap = next((g for c, g in checks if c == "total"), 0)
    subset = total_row_subset(abs(total_gap), total_row_cents) \
        if abs(total_gap) > MAX_GAP_CENTS else []
    if subset:
        refused.append(f"the total gap {total_gap:+d} cents equals printed total-row amount(s) "
                       f"{' + '.join(brl_text(c) for c in subset)}")
    if refused:
        return parser("refused as the document's: " + "; ".join(refused))
    used_all = (f"the parser used every nonzero amount the image shows ({ev['image_amounts']} on "
                f"{ev['image_pages']}, printed totals included)")
    if abs(total_gap) > MAX_GAP_CENTS:                                                  # step 4
        candidate = f"{used_all}; those amounts do not add up to the printed TOTAL GERAL"
        if conflict:
            candidate += f"; also {conflict}"
    elif conflict:
        candidate = (f"{conflict}; {used_all} and they match the TOTAL GERAL, while the per-agency "
                     f"sums disagree with the printed agency totals")
    else:
        return parser(f"not settled: {used_all} and they match the TOTAL GERAL; the per-agency "
                      f"sums disagree with the printed agency totals, which a comparison of "
                      f"amounts cannot settle, and the printed totals agree with each other")
    failing = {c: g for c, g in checks if abs(g) > MAX_GAP_CENTS}                       # step 5
    listed = REVIEWED_DOCUMENT_MONTHS.get(period)
    if listed is None:
        return parser(f"unreviewed document candidate (not in REVIEWED_DOCUMENT_MONTHS; failing "
                      f"{_gaps_text(failing)}): {candidate}")
    if failing != listed["gaps"]:
        return parser(f"unreviewed document candidate: failing {_gaps_text(failing)}, but "
                      f"reviewed on {listed['reviewed']} with {_gaps_text(listed['gaps'])}: "
                      f"{candidate}")
    return out | dict(validation_status=NOT_VALIDATED_DOCUMENT,
                      validation_reason=f"{candidate}; reviewed {listed['reviewed']}: "
                                        f"{listed['reason']}")


def build_month(period: str, files: dict[str, tuple[Path, str]]) -> tuple[dict, list[dict]]:
    rec = dict(period=period, amount_brl=np.nan, stated_total_brl=np.nan, stated_total_source="",
               total_matches=pd.NA, validation_status=NO_STATED_TOTAL, validated=False,
               validation_reason="", total_gap_cents=pd.NA,
               max_gap_cents=pd.NA, tolerance_flagged=False, tolerance_detail="", gap_detail="",
               printed_totals_conflict="", image_check="", image_pages="",
               image_amounts=pd.NA, image_unparsed=pd.NA, image_unparsed_brl="",
               image_missing=pd.NA, image_missing_brl="", image_totals_missing="",
               image_unreadable="", image_outside_amounts="", image_lines_on_total_rows="",
               n_lines=0, n_agencies=0, agencies="", status="missing_file",
               parse_method="", source_file="", source_url="", agency_totals_match=pd.NA,
               pdf_xlsx_agree=pd.NA, pdf_xlsx_resolution="", pdf_xlsx_diff_brl=np.nan,
               pdf_xlsx_lines_match=np.nan,
               title_month_mismatch=False, ambiguous_rows=0, no_payments_stated=False,
               pdf_slot_broken=False, notes="")
    docs: list[tuple[Parsed, str, Path, str]] = []
    notes: list[str] = []
    for slot in ("editable", "pdf"):
        if slot not in files:
            continue
        path, url = files[slot]
        parsed, n = parse_slot(path, slot)
        notes += n
        if slot == "pdf" and any("web page" in x for x in n):
            rec["pdf_slot_broken"] = True
        docs += [(p, slot, path, url) for p in parsed]
    readable = [d for d in docs if d[0].status == "parsed"]
    order = {"xlsx": 0, "template": 1, "pdf_in_xlsx_slot": 2, "text_pdf": 3}
    readable.sort(key=lambda d: order[d[0].method])
    if not readable:
        if docs:
            rec["status"] = "unreadable_scan"
            rec["source_file"] = docs[0][2].name
            rec["source_url"] = docs[0][3]
            rec["parse_method"] = docs[0][0].method
        rec["notes"] = "; ".join(notes)
        return rec, []

    main, slot, path, url = readable[0]
    rec.update(amount_brl=main.total, n_lines=len(main.lines),
               agencies=",".join(main.agencies), parse_method=main.method,
               source_file=f"{path.name}::{main.member}" if main.member != path.name else path.name,
               source_url=url, ambiguous_rows=main.ambiguous_rows,
               no_payments_stated=any(d[0].no_payments_stated for d in docs))
    rec["n_agencies"] = len({x["agency"] for x in main.lines})
    stated = [(d[0].stated_total, d[0].method) for d in readable if d[0].stated_total is not None]
    if main.stated_total is None and main.no_payments_stated and not main.lines:
        stated.insert(0, (0.0, main.method))
    if not stated:
        stated = [(float(sum(d[0].agency_totals.values())), f"{d[0].method}:agency_total_row")
                  for d in readable if d[0].agency_totals]
    checks: list[tuple[str, int]] = []                      # (check, gap in cents)
    if stated:
        rec["stated_total_brl"], rec["stated_total_source"] = stated[0]
        gap = main.total_cents() - cents([stated[0][0]])
        checks.append(("total", gap))
        rec["total_gap_cents"] = gap
        rec["total_matches"] = abs(gap) <= MAX_GAP_CENTS
    rec["status"] = "parsed" if stated else "no_stored_total"
    if main.agency_totals:
        own = main.agency_cents()
        per_agency = [(f"agency_total {a}", own.get(a, 0) - cents([v]))
                      for a, v in main.agency_totals.items()]
        checks += per_agency
        rec["agency_totals_match"] = all(abs(g) <= MAX_GAP_CENTS for _, g in per_agency)
    periods = {d[0].title_period for d in readable if d[0].title_period}
    rec["title_month_mismatch"] = bool(periods - {period})
    spreadsheets = [d[0] for d in readable if d[0].method in ("xlsx", "template")]
    pdfs = [d[0] for d in readable if d[0].method in ("text_pdf", "pdf_in_xlsx_slot")]
    detail_only: list[tuple[str, int]] = []
    if spreadsheets and pdfs:
        x, p = spreadsheets[0], pdfs[0]
        xc, pc = x.agency_cents(), p.agency_cents()
        pair = [(f"pdf_xlsx {a}", xc.get(a, 0) - pc.get(a, 0)) for a in sorted(set(xc) | set(pc))]
        rec["pdf_xlsx_agree"] = all(abs(g) <= MAX_GAP_CENTS for _, g in pair)
        if not rec["pdf_xlsx_agree"]:
            # The disagreement is explained when the PDF's own cells do not add up to the PDF's
            # own printed agency totals while the spreadsheet's cells do. The spreadsheet is then
            # checked against those printed totals rather than against the PDF's cells.
            pdf_self = [(f"pdf_cells_vs_own_total {a}", pc.get(a, 0) - cents([v]))
                        for a, v in p.agency_totals.items()]
            x_vs_pdf_totals = [(f"xlsx_vs_pdf_agency_total {a}", xc.get(a, 0) - cents([v]))
                               for a, v in p.agency_totals.items()]
            if p.agency_totals and any(abs(g) > MAX_GAP_CENTS for _, g in pdf_self) \
                    and all(abs(g) <= MAX_GAP_CENTS for _, g in x_vs_pdf_totals):
                rec["pdf_xlsx_resolution"] = "pdf_cells_disagree_with_own_totals"
                pair = x_vs_pdf_totals
                # The PDF's own inconsistency is reported, but it is the PDF's and the month is
                # the spreadsheet's, so it does not enter the status.
                detail_only = pdf_self
        checks += pair
        rec["pdf_xlsx_diff_brl"] = x.total - p.total
        pool = [(y["agency"], round(y["amount_brl"], 2)) for y in p.lines]
        hit = 0
        for y in x.lines:
            key = (y["agency"], round(y["amount_brl"], 2))
            if key in pool:
                pool.remove(key)
                hit += 1
        rec["pdf_xlsx_lines_match"] = hit / len(x.lines) if x.lines else 1.0
    rec.update(gap_fields(checks, bool(stated), detail_only))
    if rec["validation_status"] == NOT_VALIDATED_PARSER:
        # The image of the chosen document; for a spreadsheet month, the PDF published beside it
        # (the same table as printed, already compared with the spreadsheet agency by agency).
        image = next(((member_bytes(d[2], d[0].member), d[0]) for d in readable
                      if d[0].method in ("text_pdf", "pdf_in_xlsx_slot")), None)
        log.info("  %s fails validation (%s); reading its page images", period,
                 rec["gap_detail"])
        rec.update(explain_failure(period, checks, main.lines, main,
                                   image if image and image[0] is not None else None))
    notes += main.notes
    rec["notes"] = "; ".join(dict.fromkeys(notes))
    lines = [dict(bank_key=BANK_KEY, cnpj8=CNPJ8, period=period, **ln, basis=BASIS_AD,
                  source_file=rec["source_file"], source_url=url, parse_method=main.method)
             for ln in main.lines]
    return rec, lines


def flag_repeats(monthly: pd.DataFrame, lines: pd.DataFrame) -> pd.Series:
    blocks = {}
    for (period, agency), g in lines.groupby(["period", "agency"]):
        blocks[(period, agency)] = tuple(sorted(zip(g["category"], g["subcategory"],
                                                    g["amount_brl"].round(2))))
    periods = monthly["period"].tolist()
    out = []
    for i, period in enumerate(periods):
        rep = []
        if i:
            for (p, a), block in blocks.items():
                if p == period and block and blocks.get((periods[i - 1], a)) == block:
                    rep.append(a)
        out.append(",".join(sorted(rep)))
    return pd.Series(out, index=monthly.index)


# ---------------------------------------------------------------------------
# Sponsorship
# ---------------------------------------------------------------------------
SPONSOR_COLS = {"project": ("nome do projeto", "projeto"), "beneficiary": ("patrocinado",
                "contratado"), "cnpj": ("cnpj",), "contract_date": ("data",),
                "amount_brl": ("valor",), "justification": ("justificativa", "just"),
                "uf": ("uf",), "city": ("cidade",)}


def _sponsor_tables(member: str, b: bytes) -> list[tuple[str, pd.DataFrame]]:
    kind = kind_of(member, b)
    if kind in ("xlsx", "template"):
        wb = openpyxl.load_workbook(io.BytesIO(b), data_only=True)
        return [(ws.title, pd.DataFrame([list(r) for r in ws.iter_rows(values_only=True)]))
                for ws in wb.worksheets]
    if kind == "xls":
        return [(n, pd.DataFrame(g)) for n, g, _, _ in _grid_xlrd(b)]
    if kind == "html":
        text = b.decode("iso-8859-1")
        tables = pd.read_html(io.StringIO(text), header=None)
        return [(f"html{i}", t) for i, t in enumerate(tables)]
    raise ValueError(f"{member}: unsupported sponsorship file type {kind}")


def parse_sponsorship_file(path: Path, year: int, url: str) -> tuple[list[dict], list[str]]:
    rows, notes = [], []
    for member, b in unpack(path):
        program = "caixa_cultural" if "pocc" in fold(member) else "general"
        found = False
        for sheet, df in _sponsor_tables(member, b):
            df = df.astype(object).where(pd.notna(df), None)
            header = None
            if all(isinstance(c, str) for c in df.columns) and any(
                    "valor" in fold(c) for c in df.columns):
                df = pd.concat([pd.DataFrame([list(df.columns)], columns=df.columns), df],
                               ignore_index=True)
                df.columns = range(df.shape[1])
            for i in range(min(10, len(df))):
                cells = [fold(clean(v)) for v in df.iloc[i]]
                if any(c.startswith("valor") for c in cells) and any("projeto" in c for c in cells):
                    header = i
                    break
            if header is None:
                notes.append(f"{member}/{sheet}: no header row, sheet skipped")
                continue
            found = True
            hdr = [fold(clean(v)) for v in df.iloc[header]]
            colmap = {}
            for key, keys in SPONSOR_COLS.items():
                for j, h in enumerate(hdr):
                    if h and any(h.startswith(k) for k in keys) and j not in colmap.values():
                        colmap[key] = j
                        break
            for k in ("project", "contract_date", "amount_brl"):
                if k not in colmap:
                    raise ValueError(f"{member}/{sheet}: no {k} column in header {hdr}")
            stated = None
            for i in range(header + 1, len(df)):
                r = df.iloc[i]
                raw_amount, raw_date = r[colmap["amount_brl"]], r[colmap["contract_date"]]
                filled = [v for v in r if v not in (None, "") and clean(v)]
                if raw_date in (None, "") or isinstance(raw_date, str) and not clean(raw_date):
                    if len(filled) == 1 and isinstance(filled[0], (int, float)):
                        stated = float(filled[0])
                    elif filled and not (len(filled) == 1 and isinstance(filled[0], dt.datetime)):
                        notes.append(f"{member}/{sheet} row {i + 1}: no contract date, skipped: "
                                     f"{[clean(v)[:40] for v in filled]}")
                    continue
                amount = parse_brl(raw_amount)
                if isinstance(raw_date, (dt.datetime, dt.date)):
                    date = pd.Timestamp(raw_date)
                else:
                    date = pd.to_datetime(clean(raw_date), dayfirst=True, errors="coerce")
                rec = dict(bank_key=BANK_KEY, cnpj8=CNPJ8, year_file=year,
                           contract_date=date, program=program, amount_brl=amount,
                           basis=BASIS_SPONSOR,
                           **{k: clean(r[j]) for k, j in colmap.items()
                              if k not in ("amount_brl", "contract_date")},
                           source_file=f"{path.name}::{member}", sheet=sheet, row=i + 1,
                           source_url=url)
                rows.append(rec)
            if stated is not None:
                got = [x["amount_brl"] or 0 for x in rows
                       if x["source_file"].endswith(member) and x["sheet"] == sheet]
                gap = cents(got) - cents([stated])
                verdict = ("match" if gap == 0 else f"match within {MAX_GAP_CENTS} centavos"
                           if abs(gap) <= MAX_GAP_CENTS else "DIFFER")
                notes.append(f"{member}/{sheet}: stated total {stated:,.2f}, contracts sum "
                             f"{sum(got):,.2f} (gap {gap:+d} cents, {verdict})")
        if not found:
            raise ValueError(f"{member}: no sheet with a sponsorship header")
    return rows, notes


# ---------------------------------------------------------------------------
# Panel key
# ---------------------------------------------------------------------------
def establish_panel_key() -> str:
    codes, names = set(), {}
    for f in sorted(paths.IF_DATA_LIST.glob("*.csv")):
        d = pd.read_csv(f, dtype=str, sep=None, engine="python")
        lead = d[d["CnpjInstituicaoLider"].fillna("").str.zfill(8) == CNPJ8]
        codes |= set(lead["CodConglomeradoPrudencial"].dropna().str.strip()) - {""}
        names.update(dict(zip(d["CodInst"].str.strip(), d["NomeInstituicao"])))
    if len(codes) != 1:
        raise AssertionError(f"IF.data lists give {sorted(codes)} prudential codes for leader "
                             f"{CNPJ8}; expected exactly one")
    code = codes.pop()
    if "caixa economica federal" not in fold(names.get(code, "")):
        raise AssertionError(f"{code} is named {names.get(code)!r} in IF.data")
    cosif = pd.read_parquet(OUT_DIR / "cosif_advertising_monthly.parquet")
    rows = cosif[(cosif["level"] == "conglomerate") & (cosif["entity_key"] == code)]
    leaders = set(rows["cnpj_leader"].dropna().astype(str).str.zfill(8))
    if rows.empty or leaders != {CNPJ8}:
        raise AssertionError(f"COSIF conglomerate {code}: leader CNPJs {leaders}, expected {CNPJ8}")
    log.info("panel_key %s = %s (IF.data), confirmed by %d COSIF conglomerate rows", code,
             names[code], len(rows))
    return code


# ---------------------------------------------------------------------------
# Validation and the COSIF comparison
# ---------------------------------------------------------------------------
def validate(monthly: pd.DataFrame, lines: pd.DataFrame, contracts: pd.DataFrame,
             latest: str) -> None:
    expected = pd.period_range(pd.Period(FIRST_MONTH, "M"), pd.Period(latest, "M"), freq="M")
    expected = [p.strftime("%Y%m") for p in expected]
    got = monthly["period"].tolist()
    if got != expected or monthly["status"].isna().any():
        raise AssertionError(f"monthly grid incomplete: missing "
                             f"{sorted(set(expected) - set(got))[:10]}")                 # check 1

    # Logged before the gate below so that a failing build still reports every month's status.
    log.info("validation status (integer cents; up to %d centavos accepted and flagged):\n%s",
             MAX_GAP_CENTS, monthly["validation_status"].value_counts().to_string())
    flagged = monthly[monthly["tolerance_flagged"]]
    if len(flagged):
        log.warning("months accepted only through the %d-centavo allowance:\n%s",
                    MAX_GAP_CENTS, flagged[["period", "parse_method", "validation_status",
                                            "total_gap_cents", "tolerance_detail"]]
                    .to_string(index=False))
    document = monthly[monthly["validation_status"] == NOT_VALIDATED_DOCUMENT]
    for r in document.itertuples():
        log.warning("not validated, the DOCUMENT's inconsistency: %s (%s) total gap %s cents; "
                    "%s\n    evidence: %s", r.period, r.parse_method, r.total_gap_cents,
                    r.gap_detail, r.validation_reason)
    parser = monthly[monthly["validation_status"] == NOT_VALIDATED_PARSER]
    for r in parser.itertuples():
        log.warning("not validated, a PARSER failure: %s (%s) total gap %s cents; %s\n"
                    "    image shows, parser did not use: %s\n    evidence: %s", r.period,
                    r.parse_method, r.total_gap_cents, r.gap_detail,
                    r.image_unparsed_brl or "(none)", r.validation_reason)
    unreviewed = parser[parser["validation_reason"].str.startswith("unreviewed document candidate")]
    if len(unreviewed):
        log.warning("%d month(s) look like Caixa's inconsistency but are not in "
                    "REVIEWED_DOCUMENT_MONTHS with these gaps; they count as parser failures until "
                    "a person reviews them (module docstring, 'Reviewing a month'): %s",
                    len(unreviewed), ", ".join(f"{r.period} ({r.gap_detail})"
                                               for r in unreviewed.itertuples()))
    stale = sorted(set(REVIEWED_DOCUMENT_MONTHS) - set(document["period"]))
    if stale:
        status_of = dict(zip(monthly["period"], monthly["validation_status"]))
        log.warning("reviewed document months that did not come out as the document's this run "
                    "(their gaps or evidence changed; see their status): %s",
                    ", ".join(f"{p} {status_of.get(p, '(no such month)')}" for p in stale))

    stated = monthly[monthly["stated_total_brl"].notna()]
    log.info("check 2: %d of %d months with a stated total match it within %d centavos",
             int(stated["total_matches"].astype(bool).sum()), len(stated), MAX_GAP_CENTS)
    # The gate is there to catch a broken parser, so it counts parser failures only. A month a
    # person has reviewed as an inconsistency in Caixa's document (a typo in a printed total, a
    # cell the totals leave out) cannot be fixed by any parser change, and counting it would let
    # Caixa's typos trip a gate meant for the parser, or push the threshold down until a real
    # parser failure slipped through. Such a month is never validated; it is only kept out of
    # this count. Only reviewed months are: a parser that counts amounts it should not (a total
    # row read as data) also leaves every image amount used and the TOTAL GERAL unclosed, so an
    # unreviewed candidate counts as a parser failure.
    healthy = stated["validation_status"].isin(VALIDATED_STATUSES + (NOT_VALIDATED_DOCUMENT,))
    share = float(healthy.mean()) if len(stated) else 0.0
    log.info("check 2 (parser health): %d validated + %d document inconsistencies = %d of %d "
             "months with a stated total (%.1f%%; gate %.0f%%); %d parser failures",
             int(stated["validation_status"].isin(VALIDATED_STATUSES).sum()), len(document),
             int(healthy.sum()), len(stated), 100 * share, 100 * MIN_MATCH_SHARE, len(parser))
    if share < MIN_MATCH_SHARE:
        raise AssertionError(f"parser failures in {1 - share:.1%} of months with a stated total "
                             f"(at most {1 - MIN_MATCH_SHARE:.0%} allowed): "
                             f"{parser['period'].tolist()}")

    both = monthly[monthly["pdf_xlsx_agree"].notna()]                                  # check 3
    disagree = both[~both["pdf_xlsx_agree"].astype(bool)]
    explained = disagree[disagree["pdf_xlsx_resolution"] != ""]
    if len(explained):
        log.warning("spreadsheet and PDF disagree where the PDF's cells do not add up to its own "
                    "agency totals (spreadsheet kept):\n%s", explained[
                        ["period", "amount_brl", "pdf_xlsx_diff_brl", "pdf_xlsx_resolution"]]
                    .to_string(index=False))
    disagree = disagree[disagree["pdf_xlsx_resolution"] == ""]
    log.info("check 3: spreadsheet and PDF both parsed for %d months; %d disagree on an agency "
             "sum; line-level match median %.3f, min %.3f", len(both), len(disagree),
             both["pdf_xlsx_lines_match"].median(), both["pdf_xlsx_lines_match"].min())
    if len(disagree):
        raise AssertionError("spreadsheet and PDF disagree:\n" + disagree[
            ["period", "amount_brl", "pdf_xlsx_diff_brl", "gap_detail"]].to_string(index=False))

    if (lines["amount_brl"] < 0).any():                                                 # check 4
        raise AssertionError(f"negative amounts:\n{lines[lines['amount_brl'] < 0].head()}")

    for period, target in KNOWN_TOTALS_M.items():                                       # check 5
        v = monthly.loc[monthly["period"] == period, "amount_brl"].iloc[0] / 1e6
        if round(v, 1) != target:
            raise AssertionError(f"{period}: R$ {v:.3f}m, expected {target}m")
    log.info("check 5: April-June 2020 reproduce %s", list(KNOWN_TOTALS_M.values()))

    if contracts["amount_brl"].isna().any() or (contracts["amount_brl"] <= 0).any() \
            or contracts["contract_date"].isna().any():                                 # check 7
        bad = contracts[contracts["amount_brl"].isna() | (contracts["amount_brl"] <= 0)
                        | contracts["contract_date"].isna()]
        raise AssertionError(f"sponsorship rows without a positive amount or a date:\n"
                             f"{bad[['source_file', 'row', 'project', 'amount_brl']].head(10)}")


def print_cosif_ratios(monthly: pd.DataFrame, code: str) -> None:
    path = OUT_DIR / "cosif_advertising_quarterly.parquet"
    q = pd.read_parquet(path)
    q = q[(q["level"] == "conglomerate") & (q["entity_key"] == code) & (q["year"] == 2025)]
    m = monthly[monthly["period"].str[:4] == "2025"].copy()
    m["quarter"] = (m["period"].str[4:].astype(int) - 1) // 3 + 1
    cx = m.groupby("quarter").agg(caixa=("amount_brl", "sum"), n=("amount_brl", "count"))
    t = q.set_index("quarter")[["adv", "all3"]].join(cx)
    t.loc["H1"] = t.loc[[1, 2]].sum()
    t.loc["H2"] = t.loc[[3, 4]].sum()
    t.loc["2025"] = t.loc[[1, 2, 3, 4]].sum()
    t["ratio_adv"] = t["caixa"] / t["adv"]
    t["ratio_all3"] = t["caixa"] / t["all3"]
    for c in ("adv", "all3", "caixa"):
        t[c] = (t[c] / 1e6).round(1)
    log.info("2025: Caixa own files against COSIF conglomerate %s (R$ million; not gated):\n%s",
             code, t.round(3).to_string())
    mm = pd.read_parquet(OUT_DIR / "cosif_advertising_monthly.parquet")
    mm = mm[(mm["level"] == "conglomerate") & (mm["entity_key"] == code)
            & (mm["data_base"].str[:4] == "2025")].set_index("data_base")
    j = m.set_index("period")[["amount_brl"]].join(mm[["adv_flow", "promo_flow", "publ_flow"]])
    j["all3_flow"] = j[["adv_flow", "promo_flow", "publ_flow"]].sum(axis=1)
    log.info("2025 same-month correlation with COSIF: adv %.2f, all3 %.2f",
             j["amount_brl"].corr(j["adv_flow"]), j["amount_brl"].corr(j["all3_flow"]))


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main() -> None:
    ap = argparse.ArgumentParser(description="Caixa advertising outlays and sponsorship contracts")
    ap.add_argument("--refresh", action="store_true", help="download the lists and files again")
    ap.add_argument("--out-dir", type=Path, default=OUT_DIR)
    a = ap.parse_args()

    fetcher = Fetcher()
    ad_docs = advertising_documents(list_items(fetcher, AD_LIST, a.refresh))
    sponsor_docs = sponsorship_documents(list_items(fetcher, SPONSOR_LIST, a.refresh))
    log.info("%d monthly advertising documents and %d sponsorship files listed", len(ad_docs),
             len(sponsor_docs))
    files: dict[str, dict[str, tuple[Path, str]]] = {}
    for d in ad_docs:
        files.setdefault(d["period"], {})[d["slot"]] = (download(fetcher, d["url"], ad_path(d),
                                                                 a.refresh), d["url"])
    sponsor_files = [(download(fetcher, d["url"], sponsor_path(d), a.refresh), d)
                     for d in sponsor_docs]

    latest = max(files)
    periods = [p.strftime("%Y%m") for p in pd.period_range(pd.Period(FIRST_MONTH, "M"),
                                                           pd.Period(latest, "M"), freq="M")]
    records, all_lines = [], []
    for period in periods:
        rec, lines = build_month(period, files.get(period, {}))
        records.append(rec)
        all_lines += lines
        log.info("  %s %-16s %-16s lines %3d  sum %14s  stated %14s  %s", period, rec["status"],
                 rec["parse_method"], rec["n_lines"],
                 f"{rec['amount_brl']:,.2f}" if pd.notna(rec["amount_brl"]) else "-",
                 f"{rec['stated_total_brl']:,.2f}" if pd.notna(rec["stated_total_brl"]) else "-",
                 rec["notes"])
    monthly = pd.DataFrame(records)
    for c in ("total_gap_cents", "max_gap_cents", "image_amounts", "image_unparsed",
              "image_missing"):
        monthly[c] = monthly[c].astype("Int64")
    lines = pd.DataFrame(all_lines)
    monthly["repeated_agency_block"] = flag_repeats(monthly, lines)
    monthly.insert(0, "cnpj8", CNPJ8)
    monthly.insert(0, "bank_key", BANK_KEY)
    monthly["basis"] = BASIS_AD

    contract_rows, sponsor_notes = [], []
    for path, d in sponsor_files:
        rows, notes = parse_sponsorship_file(path, d["year"], d["url"])
        contract_rows += rows
        sponsor_notes += notes
    contracts = pd.DataFrame(contract_rows)
    contracts["date_outside_year_file"] = contracts["contract_date"].dt.year != contracts["year_file"]
    # 2022's file was converted from a PDF: some rows repeat the project, beneficiary and CNPJ
    # run together inside the text columns. Amount and date are unaffected.
    contracts["text_fields_garbled"] = contracts["project"].str.contains(
        r"\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}", regex=True)
    contracts["duplicate_row"] = contracts.duplicated(
        ["source_file", "project", "beneficiary", "contract_date", "amount_brl"], keep=False)
    for n in sponsor_notes:
        log.info("  sponsorship: %s", n)

    code = establish_panel_key()                                                        # check 6
    for frame in (monthly, lines, contracts):
        frame["panel_key"] = code
    validate(monthly, lines, contracts, latest)

    a.out_dir.mkdir(parents=True, exist_ok=True)
    for name, frame in (("caixa_advertising_lines", lines),
                        ("caixa_advertising_monthly", monthly),
                        ("caixa_sponsorship_contracts", contracts)):
        frame.to_parquet(a.out_dir / f"{name}.parquet", index=False)
        frame.to_csv(a.out_dir / f"{name}.csv", index=False)
        log.info("%s: %d rows -> %s", name, len(frame), a.out_dir / f"{name}.parquet")

    monthly["year"] = monthly["period"].str[:4]
    cov = (monthly.groupby(["year", "status"]).size().unstack(fill_value=0)
           .join(monthly.groupby("year")["amount_brl"].sum().div(1e6).round(1)
                 .rename("sum_brl_m")))
    log.info("coverage by year:\n%s", cov.to_string())
    flags = monthly[(monthly["repeated_agency_block"] != "") | monthly["title_month_mismatch"]
                    | (monthly["ambiguous_rows"] > 0) | monthly["no_payments_stated"]
                    | monthly["pdf_slot_broken"] | (monthly["agency_totals_match"] == False)  # noqa: E712
                    | monthly["tolerance_flagged"]]
    log.info("flagged months:\n%s", flags[["period", "parse_method", "repeated_agency_block",
                                           "title_month_mismatch", "ambiguous_rows",
                                           "no_payments_stated", "pdf_slot_broken",
                                           "agency_totals_match", "tolerance_flagged"]]
             .to_string(index=False))
    sp = (contracts.groupby(["year_file", "program"])
          .agg(contracts=("amount_brl", "size"), sum_brl_m=("amount_brl", "sum"),
               outside_year=("date_outside_year_file", "sum")))
    sp["sum_brl_m"] = (sp["sum_brl_m"] / 1e6).round(1)
    log.info("sponsorship contracts by file year:\n%s", sp.to_string())
    print_cosif_ratios(monthly, code)


if __name__ == "__main__":
    main()
